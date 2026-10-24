"""Validate canonical data and derive static data + departmental Excel workbooks."""
import json,pathlib,sys,hashlib,re
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from deterministic_xlsx import save_deterministic_xlsx
from public_metadata import check_public_tree
ROOT=pathlib.Path(__file__).resolve().parents[1]
def valid_url(v):return isinstance(v,str) and urlparse(v).scheme in ('http','https') and bool(urlparse(v).netloc)
def mentor_official_url(mentor):
 # Explicit false flags veto shortcuts without weakening verified employment evidence.
 if not mentor or mentor.get('employment_status')!='active_officially_listed' or mentor.get('quick_link_eligible') is False or mentor.get('navigation_verified') is False:return ''
 url=mentor.get('official_url')
 try:
  if not valid_url(url):return ''
  parsed=urlparse(url)
  if not parsed.hostname or any(c.isspace() for c in parsed.hostname):return ''
  parsed.port  # Access validates the numeric port and range, like the browser URL parser.
  return url
 except ValueError:return ''
def mentor_navigation_notice(mentor):
 if not mentor or mentor.get('employment_status')!='active_officially_listed' or mentor_official_url(mentor):return ''
 reason=mentor.get('profile_navigation_unavailable_reason')
 return '任职已核实；'+((reason.strip() if isinstance(reason,str) else '') or '导师主页访问未核实，暂不提供快捷链接。')

def safe_id(value):return isinstance(value,str) and bool(re.fullmatch(r'[A-Za-z0-9_-]+',value))
def school_index(d):
 schools={row['id']:row for row in d['institutions']}
 assert len(schools)==len(d['institutions']), 'Duplicate school ID'
 assert all(safe_id(sid) for sid in schools), 'Unsafe school ID'
 assert all(valid_url(row.get('official_url')) and row.get('official_name') for row in schools.values()), 'School identity needs a name and official URL'
 return schools

def validate_school_scope(d):
 schools=school_index(d)
 campuses={c['id']:c for c in d['campuses']}
 assert len(campuses)==len(d['campuses']), 'Duplicate campus ID'
 orgs={o['id']:o for o in d['org_units']}
 assert len(orgs)==len(d['org_units']), 'Duplicate organization ID'
 assert all(safe_id(oid) for oid in orgs), 'Unsafe organization ID'
 legacy={'harbin','weihai','shenzhen'}
 for c in campuses.values():
  assert c.get('school_id') in schools, 'Unknown campus school'
  assert safe_id(c['id']), 'Unsafe campus ID'
  assert (c['school_id']=='hit' and c['id'] in legacy) or c['id'].startswith(c['school_id']+'-'), 'New campus IDs must be school-prefixed'
 for key in ['org_units','people','pending_people']:
  for row in d.get(key,[]):
   sid=row.get('school_id');assert sid in schools, 'Unknown record school'
   assert isinstance(row.get('campus_ids'),list), 'Campus associations must be explicit; empty is allowed'
   for cid in row['campus_ids']:
    assert cid in campuses and campuses[cid]['school_id']==sid, 'Cross-school or unknown campus reference'
   if key=='org_units':
    assert row.get('name') and row.get('type') and valid_url(row.get('source_url')), 'Organization needs a name, type and source URL'
    parent=row.get('parent_id')
    if parent in orgs:assert orgs[parent]['school_id']==sid, 'Cross-school organization parent'
    elif parent:
     # Preserve historical HIT parent aliases without creating synthetic org nodes.
     legacy_parent=sid=='hit' and parent in {'hit-'+c for c in legacy}
     assert parent==sid or (parent in campuses and campuses[parent]['school_id']==sid) or legacy_parent, 'Unknown or cross-school organization parent'
   else:
    oid=row.get('department_id') or row.get('organization_id')
    assert oid in orgs and orgs[oid]['school_id']==sid, 'Cross-school or unknown person department'
 return True

def validate(d):
 validate_school_scope(d)
 pending=d.get('pending_people',[])
 assert all(p.get('profile_status') in ('roster_identity_pending','profile_details_pending') and not p.get('education') and not p.get('mentors') and not p.get('publications') and not p.get('professional_research_advisors') for p in pending),'Pending identities must stay outside confirmed collections'
 assert not ({p['person_id'] for p in pending}&{p['person_id'] for p in d['people']}),'Pending identity duplicated as detailed profile'
 ids=[u['id'] for u in d['org_units']]; assert len(ids)==len(set(ids)), 'Duplicate organization ID'
 people=d['people']; assert len({p['person_id'] for p in people})==len(people),'Duplicate person ID'
 for p in people:
  assert p['department_id'] in ids,'Unknown department'
  assert valid_url(p['official_url']),'Unsafe profile URL'
  teaching=p.get('teaching')
  assert teaching is None or isinstance(teaching,list),'Teaching must be an array or null'
  for fact in teaching or []:
   assert isinstance(fact,dict),'Teaching facts must be objects'
   assert (isinstance(fact.get('name'),str) and fact['name'].strip()) or (isinstance(fact.get('course'),str) and fact['course'].strip()) or (isinstance(fact.get('courses'),list) and fact['courses'] and all(isinstance(v,str) and v.strip() for v in fact['courses'])),'Teaching facts require a named topic or legacy course list'
   assert isinstance(fact.get('source_id'),str) and fact['source_id'].strip(),'Teaching facts require a source ID'
   assert fact.get('role') is None or isinstance(fact['role'],str) and fact['role'].strip(),'Teaching role must be nonempty when provided'
   assert valid_url(fact.get('source_url')),'Teaching facts require a safe evidence URL'
   sources=[s for s in p['sources'] if s.get('id')==fact['source_id']]
   assert len(sources)==1,'Teaching references unknown or ambiguous source'
   assert sources[0].get('url')==fact['source_url'],'Teaching source ID and URL must agree'
   assert all(fact.get(key) is None or isinstance(fact[key],str) for key in ['start','note']),'Teaching time and note must be text or null'
   assert not any(fact.get(key) for key in ['education_event_id','education_event_ids','appointment_event_id']),'Teaching must not bind education or appointment events'
   assert not any(fact.get(key) is True for key in ['degree_awarded','include_in_mentorship_graph','include_in_verified_mentor_count','include_in_education_count','include_in_appointment_count','include_in_publication_count']),'Teaching must stay outside degree and mentor claims'
  for collection in ['education','mentors','appointments','publications','committee_members']:
   for row in p.get(collection,[]):
    assert valid_url(row.get('source_url')),'Missing evidence URL'
  for e in p['education']:
   assert e['institution_985'] in ('yes','no','unknown')
   assert e['region'] in ('mainland','hmt','overseas','unknown')
   assert 'event_kind' in e and e.get('degree_awarded') in (True,False,None),'Missing event/award semantics'
   assert e.get('degree_awarded') is not True or e['event_kind']=='degree','Only degree events may confirm an award'
  event_ids=[e['event_id'] for e in p['education'] if e.get('event_id')]
  assert len(event_ids)==len(set(event_ids)),'Duplicate education event ID'
  for c in p.get('research_collaborators',[]):
   assert c.get('relationship_kind')!='professional_research_advisor','Explicit professional guidance must not be reduced to ordinary collaboration'
   assert c.get('include_in_mentorship_graph') is False,'Collaborator must not be promoted to mentorship'
   if c.get('relationship_kind')=='visiting_research_collaborator':
    assert c.get('education_event_id') in event_ids,'Collaborator references unknown education event'
    event=next(e for e in p['education'] if e.get('event_id')==c['education_event_id'])
    assert event.get('degree_awarded') is False,'Visiting collaborator must bind a non-award event'
   if c.get('appointment_event_id'):
    assert any(a.get('event_id')==c['appointment_event_id'] for a in p['appointments']),'Collaborator references unknown appointment'
   assert not(c.get('education_event_id') and c.get('appointment_event_id')),'Collaborator cannot bind two event kinds'
   assert valid_url(c.get('source_url')),'Collaborator requires evidence'

  assert isinstance(p.get('professional_research_advisors',[]),list),'Professional guidance must be an array'
  for c in p.get('professional_research_advisors',[]):
   assert isinstance(c,dict),'Professional guidance rows must be objects'
   assert c.get('relationship_kind')=='professional_research_advisor','Professional guidance must preserve its verified relationship kind'
   assert c.get('relationship_status')=='verified_direct_statement','Professional guidance requires a verified direct relationship'
   assert c.get('include_in_mentorship_graph') is False and c.get('include_in_verified_mentor_count') is False,'Independent professional guidance must stay outside main mentor counts and graph'
   assert c.get('include_in_degree_mentorship_graph') is False,'Professional guidance must stay outside degree mentorship'
   assert c.get('include_in_verified_professional_guidance_count') is True,'Verified professional guidance count must be explicit'
   assert 'education_event_id' in c,'Professional guidance must explicitly have no degree binding or a verified non-award research-event binding'
   assert not c.get('education_event_ids'),'Professional guidance cannot use an ambiguous plural event binding'
   if c['education_event_id'] is not None:
    assert professional_guidance_event(p,c),'Professional guidance may bind only a known explicit non-award professional research event'
    assert not c.get('appointment_event_id'),'Professional guidance cannot bind two event kinds'
   assert isinstance(c.get('name'),str) and c['name'].strip() and valid_url(c.get('source_url')),'Professional guidance requires direct source URL and name'
   if c.get('appointment_event_id') is not None:
    assert isinstance(c['appointment_event_id'],str) and c['appointment_event_id'].strip(),'Professional guidance appointment ID cannot be empty'
    assert any(a.get('event_id')==c['appointment_event_id'] for a in p['appointments']),'Professional guidance references unknown appointment'
   sources={row.get('id'):row.get('url') for row in p['sources']}
   for key,urlkey in [('source_id','source_url'),('profile_source_id','profile_source_url')]:
    if c.get(key) is not None:
     assert c[key] in sources,'Professional guidance references unknown source'
     assert sources[c[key]]==c.get(urlkey),'Professional guidance source ID and URL must agree'
   if 'source_ids' in c:
    assert isinstance(c['source_ids'],list) and all(isinstance(s,str) and s in sources for s in c['source_ids']),'Professional guidance references unknown evidence sources'
   if c.get('employment_status')=='active_officially_listed':
    assert valid_url(c.get('profile_source_url')),'Active professional guidance needs public employment evidence'
    unavailable=(c.get('official_url') is None and c.get('quick_link_eligible') is False and c.get('navigation_verified') is False and isinstance(c.get('profile_navigation_unavailable_reason'),str) and bool(c['profile_navigation_unavailable_reason'].strip()))
    assert valid_url(c.get('official_url')) or unavailable,'Active professional guidance needs homepage or explicit verified navigation limitation'

  for c in p.get('project_guidance',[]):
   assert c.get('include_in_mentorship_graph') is False,'Project guidance must stay outside the mentorship graph'
   assert c.get('include_in_verified_mentor_count') is False,'Project guidance must stay outside mentor counts'
   assert c.get('education_event_id') is None,'Project guidance must not imply a degree-event binding'
   assert c.get('name') and valid_url(c.get('source_url')),'Project guidance requires a name and direct evidence URL'
   source_ids={row.get('id') for row in p['sources']}
   for key in ['source_id','profile_source_id']:
    assert not c.get(key) or c[key] in source_ids,'Project guidance references unknown source'

  for m in p['mentors']:
   assert m.get('type')!='research_project_advisor' and m.get('relationship_kind')!='research_project_advisor','Project guidance must stay separate from mentors'
   assert not m.get('education_event_id') or m['education_event_id'] in event_ids,'Mentor references unknown education event'
   if 'education_event_ids' in m:
    refs=m['education_event_ids']
    assert isinstance(refs,list) and refs and all(isinstance(v,str) and v.strip() for v in refs),'Mentor event IDs must be a nonempty array of nonempty strings'
    assert len(refs)==len(set(refs)),'Duplicate mentor education event ID'
    assert all(v in event_ids for v in refs),'Mentor references unknown education events'
    assert not m.get('education_event_id') or refs==[m['education_event_id']],'Mentor singular and plural event bindings conflict'
    if len(refs)>1:
     assert all(isinstance(m.get(k),str) and m[k].strip() for k in ['stage_scope','education_stage']),'Combined mentorship must preserve its explicit stage scope'
   if m.get('relationship_kind')=='professional_research_advisor':
    # Preserve reviewed legacy non-award event bindings, never invent one for the new independent collection.
    event=next((e for e in p['education'] if m.get('education_event_id') and e.get('event_id')==m['education_event_id']),None)
    assert event and event.get('degree_awarded') is False and m.get('include_in_degree_mentorship_graph') is not False,'Unbound professional guidance must stay in its independent collection'
   if m['employment_status']=='active_officially_listed':
    assert valid_url(m.get('profile_source_url')),'Active mentor needs public employment evidence'
    unavailable = (m.get('official_url') is None and m.get('quick_link_eligible') is False and m.get('navigation_verified') is False and isinstance(m.get('profile_navigation_unavailable_reason'), str) and bool(m['profile_navigation_unavailable_reason'].strip()))
    assert valid_url(m.get('official_url')) or unavailable,'Active mentor needs homepage or explicit verified navigation limitation'
 for u in d['org_units']:
  seen=set();cur=u
  while cur['parent_id'] in ids:
   assert cur['id'] not in seen,'Organization cycle';seen.add(cur['id']);cur=next(v for v in d['org_units'] if v['id']==cur['parent_id'])
 return True
def clean(v):
 if isinstance(v,(list,dict)):v=json.dumps(v,ensure_ascii=False)
 if v is None:return ''
 if isinstance(v,str) and v and v[:1] in '=+-@\t\r':return "'"+v
 return v
def sheet(wb,name,rows,disabled_sources_by_person=None):
 ws=wb.create_sheet(name)
 for row in rows:ws.append([clean(v) for v in row])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='102A42');cell.font=Font(color='FFFFFF',bold=True)
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if valid_url(cell.value) and cell.value not in (disabled_sources_by_person or {}).get(ws.cell(cell.row,1).value,set()):cell.hyperlink=cell.value;cell.font=Font(color='078B92',underline='single')
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col)*1.4))
 return ws
def disabled_sources_by_person(people):
 # Only explicit source metadata for the person named on the workbook row controls links.
 result={}
 for p in people:
  blocked={s['url'] for s in p.get('sources',[]) if s.get('navigation_enabled') is False or s.get('navigation_available') is False}
  for rows in p.values():
   if isinstance(rows,list):blocked.update(row['source_url'] for row in rows if isinstance(row,dict) and row.get('source_navigation_enabled') is False and valid_url(row.get('source_url')))
  if blocked:assert sum(x['name']==p['name'] for x in people)==1,'Ambiguous same-name source navigation; retain person separation before exporting'
  result.setdefault(p['name'],set()).update(blocked)
 return result

def source_availability_note(person,url):
 source=next((s for s in person.get('sources',[]) if s.get('url')==url and (s.get('navigation_enabled') is False or s.get('navigation_available') is False)),None)
 return source_evidence_text(source) if source else ''

def source_evidence_text(source):
 label=source.get('evidence_label') or source.get('evidence_status_label')
 status=source.get('evidence_status')
 return '；'.join(str(v) for v in [label,status] if v) if label and label!=status else status or label

def publication_value(publication,key):
 raw=publication.get(key);label=publication.get(key+'_label') if key in ('type','link_status') else None
 return '；'.join(str(v) for v in [label,raw] if v) if label and label!=raw else raw if raw is not None else label

def profile_availability_label(person):
 if person.get('official_profile_availability_label'):return person['official_profile_availability_label']
 status='个人主页当前不可用' if person.get('official_profile_availability')=='unavailable' or person.get('official_profile_navigation_eligible') is False else '公开个人页（访问与内容时效见来源记录）'
 return '；'.join(v for v in [person.get('official_profile_link_label'),status] if v)

def teaching_summary(person):
 # Teaching is evidence about the teacher's work, never the teacher's own education.
 rows=person.get('teaching')
 if not isinstance(rows,list) or not rows:return ''
 facts=[]
 for row in rows:
  if not isinstance(row,dict):continue
  facts.append('；'.join(str(v) for v in [
   row.get('name') or row.get('course') or '、'.join(row.get('courses') or []) or '教学主题未披露',row.get('role') or '角色未披露',
   '教学／课程开设时间：'+str(row.get('start') or '时间未披露'),row.get('note'),
   '来源ID：'+str(row.get('source_id') or '未披露'),row.get('source_url'),source_availability_note(person,row.get('source_url'))
  ] if v))
 if not facts:return ''
 return '公开教学事实（非本人求学学历；教学／课程开设时间不是入学、授位或任职起点；不计入教育、任职、论文、导师统计及师承图）：'+' | '.join(facts)

def professional_guidance_event(person,guidance):
 event_id=guidance.get('education_event_id')
 if not isinstance(event_id,str) or not event_id.strip():return {}
 events=[e for e in person.get('education',[]) if e.get('event_id')==event_id]
 return events[0] if len(events)==1 and events[0].get('event_kind')=='professional_research' and events[0].get('degree_awarded') is False else {}

def combined_mentor_note(mentor):
 if not isinstance(mentor.get('education_event_ids'),list):return ''
 return '；'.join(str(v) for v in [
  '合述仅计一条关系，不拆分阶段导师认定' if len(mentor['education_event_ids'])>1 else '明确事件绑定',
  mentor.get('stage_scope'),mentor.get('relationship_scope'),mentor.get('scope_label'),
  '关联经历ID：'+json.dumps(mentor['education_event_ids'],ensure_ascii=False),
  ('关系明确机构：'+str(mentor['institution'])) if mentor.get('institution') else '合述机构未单独披露',
  '指导起止：'+str(mentor.get('start') or '起始时间未披露')+' 至 '+str(mentor.get('end') or '结束时间未披露'),
  mentor.get('note'),mentor.get('evidence_note')
 ] if v)

def professional_event_note(person,guidance):
 event=professional_guidance_event(person,guidance)
 if not event:return ''
 return '对应非学位科研经历：'+str(event['event_id'])+'；'+str(event.get('institution') or '机构未披露')+'；'+str(event.get('start') or '起始时间未披露')+' 至 '+str(event.get('end') or '结束时间未披露')+'（经历区间，不等于指导起止；非博士后认定）'

def source_retrieval_date_summary(person):
 # Retrieval dates are source metadata, not a whole-person review or current-employment check.
 sources=person.get('sources') if isinstance(person.get('sources'),list) else []
 dates=set();missing=not sources;invalid=False
 for source in sources:
  value=source.get('retrieved_at') if isinstance(source,dict) else None
  if value is None or isinstance(value,str) and not value.strip():missing=True;continue
  if not isinstance(value,str) or not re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}',value):invalid=True;continue
  year,month,day=map(int,value.split('-'));leap=year%4==0 and (year%100!=0 or year%400==0)
  days=[31,29 if leap else 28,31,30,31,30,31,31,30,31,30,31]
  if year<1 or not 1<=month<=12 or not 1<=day<=days[month-1]:invalid=True;continue
  dates.add(value)
 ordered=sorted(dates);summary=(ordered[0]+' 至 '+ordered[-1]) if len(ordered)>1 else ordered[0] if ordered else '未知'
 notes=[]
 if missing:notes.append('部分来源日期未单独记录' if ordered else '来源日期未单独记录')
 if invalid:notes.append('部分来源日期异常' if ordered else '来源日期异常')
 return summary+('（'+'；'.join(notes)+'）' if notes else '')

def department_tables(people,updated_at,institutions=None):
 # Retain the existing call signature; the dataset update date is not a retrieval/review date.
 schools={p.get("school_id") for p in people}
 assert len(schools)<=1 and None not in schools, "Department export cannot mix schools or omit school IDs"
 departments={p.get("department_id") for p in people}
 assert len(departments)<=1 and all(safe_id(dept) for dept in departments), "Department export cannot mix departments or omit department IDs; group people by department_id"
 school_id=next(iter(schools),"")
 names={s["id"]:s["official_name"] for s in (institutions or [])}
 school_name=names.get(school_id,school_id)
 tables=[('总览',[['姓名','校区','院系','官方主页','已记录来源检索日期','缺项与冲突','范围','主页可用性','替代官方职业报道（非个人主页）'],*[[p['name'],p.get('campus') or '校区未标注',p['department'],'' if p.get('official_profile_availability')=='unavailable' or p.get('official_profile_navigation_eligible') is False else p['official_url'],source_retrieval_date_summary(p),'；'.join(p['unknowns']),'分批核实资料，不是完整名录',profile_availability_label(p),'；'.join(str(s.get('title','官方职业报道'))+'：'+str(s.get('url','')) for s in p.get('alternative_professional_sources',[]))] for p in people]])]
 # Append inside the existing overview scope cell: keep all five sheets and every column.
 for p,row in zip(people,tables[0][1][1:]):
  summary=teaching_summary(p)
  current=p.get('current_employment') or {}
  affiliation='；'.join(str(v) for v in [p.get('sampling_affiliation_label'),'已核现任单位：'+str(current['institution']) if current.get('institution') else None,current.get('note'),current.get('source_url')] if v) if p.get('sampling_affiliation_label') else ''
  if summary or affiliation:row[6]='；'.join(v for v in [row[6],p.get('profile_coverage_label'),affiliation,summary] if v)
  row[6]+='；人物整体核验日期未单独记录'
 # Explicit source labels, not URL hosts, opt this table into neutral headings. Positions stay fixed.
 if any(p.get('official_profile_link_label') or p.get('official_profile_availability_label') for p in people):
  tables[0][1][0][3]='个人主页／职业网站（来源类型见说明）'
  tables[0][1][0][8]='替代职业资料来源（非已核个人主页）'
 fields=['event_kind','degree','institution','start','end','field','country','source_region','continent','institution_985','degree_awarded','degree_awarded_date','degree_award_source_url','note','geography_basis','source_url']
 rows=[['姓名','事件类型','阶段/职位','机构','开始','结束','专业','国家','地区','洲','历史985分类','确证授位','授位日期','授位来源','备注及冲突','地域口径','来源']]
 for p in people:
  events=p['education']+[dict(a,event_kind='employment',degree=a.get('title')) for a in p['appointments']]
  rows.extend([[p['name'],*[str(e.get(k) or '')+'（层次暂对应，未核实）' if k=='degree' and str(e.get('field_status',{}).get('degree','')).startswith('provisional_') else e.get(k) for k in fields]] for e in events])
  for c in p.get('research_collaborators',[]):
   event_id=c.get('education_event_id') or c.get('appointment_event_id')
   events=p['education'] if c.get('education_event_id') else p['appointments']
   e=next((e for e in events if event_id and e.get('event_id')==event_id), {})
   row=dict(event_kind=c.get('relationship_kind','research_collaborator'),degree=c.get('type'),institution=c.get('institution'),start=e.get('start'),end=e.get('end'),source_url=c.get('source_url'),note='；'.join(str(v) for v in [c.get('name'),c.get('description'),c.get('employment_status_note'),c.get('note') or c.get('evidence_note'),'非导师认定，不计师承',c.get('profile_source_url')] if v))
   rows.append([p['name'],*[row.get(k) for k in fields]])
  for c in p.get('project_guidance',[]):
   # Project context is not a degree event. Never join null/missing education IDs.
   row=dict(event_kind='project_guidance',degree='项目指导（非学位师承）',institution=c.get('institution'),start=c.get('start'),end=c.get('end'),degree_awarded=False,source_url=c.get('source_url'),note='；'.join(str(v) for v in [c.get('name'),c.get('type_label') or c.get('type'),c.get('project_name'),c.get('education_stage'),c.get('relationship_scope'),c.get('description'),c.get('employment_status_note'),c.get('note'),c.get('scope_label'),'非学位师承，不计入导师统计或师承关系图',c.get('profile_source_url')] if v))
   rows.append([p['name'],*[row.get(k) for k in fields]])
 tables.append(('经历',rows))
 mentor_fields=['name','type','education_stage','description','employment_status','employment_status_note','official_url','source_url','profile_source_url']
 rows=[['姓名','导师姓名','关系类型','阶段','简介','在职状态','状态说明','已核在职主页','关系来源','任职来源']]
 for p in people:
  mentors=[dict(m,type=m.get('type_label') or m.get('type'),official_url=mentor_official_url(m),employment_status_note='；'.join(v for v in [m.get('employment_status_note'),mentor_navigation_notice(m),combined_mentor_note(m)] if v) if mentor_navigation_notice(m) or combined_mentor_note(m) else m.get('employment_status_note')) for m in p['mentors']]+[dict(c,type=c.get('label','博士指导委员会成员（非师承）'),official_url='') for c in p.get('committee_members',[])]
  # Independent occupational guidance is preserved, not silently converted to a degree mentor.
  mentors += [dict(c,type='职业科研指导（非学位）',education_stage=c.get('education_stage') or '职业科研阶段（非学历绑定）',official_url=mentor_official_url(c),employment_status_note='；'.join(str(v) for v in [c.get('type_label'),c.get('type'),c.get('relationship_kind'),('机构：'+str(c['institution'])) if c.get('institution') else None,('指导起止：'+str(c.get('start') or '起始时间未披露')+' 至 '+str(c.get('end') or '结束时间未披露')) if c.get('start') or c.get('end') or professional_guidance_event(p,c) else None,professional_event_note(p,c),c.get('employment_status_note'),c.get('relationship_scope'),c.get('note'),c.get('evidence_note'),c.get('scope_label'),mentor_navigation_notice(c),'独立职业指导，不计入学位师承图；已核职业指导计数：'+('1' if c.get('include_in_verified_professional_guidance_count') is True else '0')] if v)) for c in p.get('professional_research_advisors',[])]
  rows.extend([[p['name'],*[m.get(k) for k in mentor_fields]] for m in mentors])
 tables.append(('导师',rows))
 for label,key,fields,headers in [('论文','publications',['title','year','doi','url','type','source_url','identity_evidence','link_status','note'],['标题','年份','DOI','链接','类型','收录来源','作者归属依据','链接核验范围','备注']),('来源','sources',['id','url','evidence_status','retrieved_at'],['来源ID','链接','核验方式','查阅日期'])]:
  rows=[['姓名',*headers],*[[p['name'],*[source_evidence_text(r) if k=='evidence_status' else publication_value(r,k) if key=='publications' else r.get(k) for k in fields]] for p in people for r in p[key]]]
  if key=='sources':
   for p in people:
    seen={row.get('url') for row in p['sources']}
    for e in p['education']:
     url=e.get('classification_source')
     if e.get('institution_985') in ('yes','no') and valid_url(url) and url not in seen:
      rows.append([p['name'],'institution-classification-'+str(len(seen)),url,'历史985院校分类参考；非个人授位或就读当年建设身份的证明','未单独记录']);seen.add(url)
  tables.append((label,rows))
 for _,rows in tables:
  rows[0].extend(['所属学校ID','所属学校'])
  for row in rows[1:]:row.extend([school_id,school_name])
 return tables

def department_exports(d):
 validate_school_scope(d)
 units={u['id']:u for u in d['org_units']}
 names={s['id']:s['official_name'] for s in d['institutions']}
 result=[]
 for dept in sorted({p['department_id'] for p in d['people']}):
  school_id=units[dept]['school_id']
  people=[p for p in d['people'] if p['department_id']==dept and p['school_id']==school_id]
  result.append({'department_id':dept,'school_id':school_id,'school_name':names[school_id],'path':'exports/'+dept+'.xlsx','people_count':len(people)})
 return result

def build(generate_excel=True):
 check_public_tree(ROOT,phase='preflight')
 d=json.loads((ROOT/'data/atlas.json').read_text());validate(d);exports=department_exports(d)
 d['build_id']=hashlib.sha256((ROOT/'data/atlas.json').read_bytes()+(ROOT/'data/rosters.json').read_bytes()+pathlib.Path(__file__).read_bytes()+(ROOT/'scripts/build_rosters.py').read_bytes()+(ROOT/'scripts/deterministic_xlsx.py').read_bytes()).hexdigest()[:16]
 out=ROOT/'site/exports';out.mkdir(exist_ok=True)
 for export in exports:
  dept=export['department_id'];people=[p for p in d['people'] if p['department_id']==dept and p['school_id']==export['school_id']];path=out/(dept+'.xlsx')
  if generate_excel:
   wb=Workbook();wb.remove(wb.active)
   for name,rows in department_tables(people,d['updated_at'],d['institutions']):sheet(wb,name,rows,disabled_sources_by_person(people))
   save_deterministic_xlsx(wb,path)
 from build_rosters import build_rosters
 build_rosters(d,generate_excel=generate_excel)
 roster_path=ROOT/'site/rosters.json';roster_snapshot=json.loads(roster_path.read_text());roster_snapshot['build_id']=d['build_id'];roster_path.write_text(json.dumps(roster_snapshot,ensure_ascii=False,indent=2)+'\n')
 d['exports']=exports;(ROOT/'site/atlas.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
 check_public_tree(ROOT,phase='final',include_workbooks=generate_excel)
 print(f'Validated {len(d["people"])} detailed profiles; prepared {len(exports)} department exports and website snapshot')
 print('Public source metadata: final JSON/CSV/XLSX checked' if generate_excel else 'Public source metadata: JSON/CSV checked; XLSX not regenerated or checked (--data-only)')
if __name__=='__main__':
 import argparse
 parser=argparse.ArgumentParser();parser.add_argument('--data-only',action='store_true');args=parser.parse_args();build(generate_excel=not args.data_only)
