from atlas_test_io import load_atlas
from published_fixture import published_root
from source_date_history import historical_bytes
"""Independent release23 NJU/HNU source-pinned, offline regression.

Run from a repository tests directory or use ACADEMIC_ATLAS_TEST_ROOT. Runtime
requires only the release repository and its normal test dependencies, never the
research directory, source packages, normalization code, immutable baseline, or
network. Original public leaf fingerprints and shape manifests were derived from
the two SHA-pinned raw source packages, not the candidate or normalization output.
Only the non-factual public_raw_cache:false marker was omitted from source shapes.
All array lengths, original values, conflicts, precision, unknowns, and provenance
are pinned. Documented storage aliases are reversed without repairing facts.
Canonical tables and actual JS consumers run in memory and do not write builds.
This checks the approved selected sample, not live sources or whole-school staff.
"""
from release_history import release25_view
import copy
import hashlib
import json
import os
import pathlib
import subprocess
import sys
import unittest
from collections import Counter
from unittest import mock

sys.dont_write_bytecode = True
ROOT = pathlib.Path(os.environ.get('ACADEMIC_ATLAS_TEST_ROOT', pathlib.Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(ROOT / 'scripts'))
import build
from build_rosters import validate_rosters


def object_sha(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest()

SOURCE_PACKAGE_SHA256 = {'nju': '141e34961e4b39300cd05e1e32fe4de0af5676e20601153d0e1d6d8ff50fe471',
 'hnu': '7ad107b353fa4da5a549cfbe89e14fbb3ec6500014d0c6d782abda41b5253a54'}

REVIEW_SHA256 = {'nju': '1b7966ee9adeeb3db35592723f5db4c1d3ca12bfde62655ccd4612d4122e4147',
 'hnu': '746275376c7b6df96457535630128e0923246417a9ff43a0841ae5d3093a83b7'}

BASELINE_ATLAS_PREFIXES = {'limitations': (6, 'a49996988d143a02303ec539cb493667cdc9b28a287271ae141f6cfd7b95021a'),
 'campuses': (4, 'a29d9956fbbc284b9cdafc640ac9057dd4512c2181d89780a92a6a8bf01e6505'),
 'org_units': (218, 'e466d8b7764c4e1baf2ca17059473938a0c641fc118dded67f32f100210d2052'),
 'people': (535, '192b96b35f63bb1d029007c95146a3f9819a5c8906d946189d3ac0ec6d7274ab'),
 'profile_batches': (130, 'e22423c9977fe39f47249d8b02895dee1a3df98d406a0bcaa063de7324145b80'),
 'pending_people': (83, '7bf2a61c24b2737d8d175bc65545de28d54f65214313881144efaa43dcc89bb5'),
 'normalization_evidence': (12, '697486b851572c011c9365733811793ca59c1af3cc3c77395badc8d11b772a35'),
 'institutions': (39, '8234ea78331d931b8a54c5f40d4427a7844e7fcd621d0832a69e5f2df1b41848')}

BASELINE_ATLAS_UNCHANGED = {'schema_version': '5a4747fcdafec66d8d8812022e3b7d2666d886ec0e0ee7926ec8c50db40f29d5',
 'institution': '9b9fab8a1164afafcf855e8d6ccea17a09e973b8ea2fd8a4831ad681490e8d01',
 'coverage': '8301d88ec52e91fd40964298ce12e179dc0c4ca513b3a0cefa9ad57db895f597'}

BASELINE_ROSTER_PREFIXES = {'summary': (137, '5110eac4decb7961303b9288e6b300dfb165fcc3edc05e2f1382611b801b1c18'),
 'coverage': (208, '32d8d226fb443bcf1472360e2728022cde072d66aad96f94594d2909cdc15834'),
 'entries': (1700, 'f80e99fc35fc648ba1ce96f976d8569c804b5897fbae5bad518e6080fc5db0a7'),
 'conflicts': (30, '0f809e3a2c4ef1857b03479230f35a342b703ccf5f8a99ef032271d150c44c09'),
 'limitations': (35, '4fe8dea8da59fb3789194955a1b22aac4453cbf0e6a4d1cc2e82d172ba77c116'),
 'dedup_candidates': (60, '0a1adbff0d7b0ae2fe42cceff9a89e5e9eb6d164044456f319cfab6c9d1b1653'),
 'exclusions': (3, '89dc768c034efe9de60bb4199c92be749ea7f1bb1405daceff0e3ca88fa5228d'),
 'pending_checks': (2, '143ba1254ce843007ed487ca7f7ef2c684912b54c0580c6bb3f7157b7bc61048'),
 'source_snapshots': (36, 'bae26c4aa31edf9ee3212a7723508241f475441863f03f69a8058264ad3b81c2'),
 'batches': (103, 'b13c074d4c77f623abc6c11ced7dbd83aca2e2842830f7d8252d1a8fe1eede28'),
 'source_observations': (1738, '579a7202c8ff2a4372e8248e13b83bfc0b1f60fd5ab5c6f0f52f122c5af3c166'),
 'consolidations': (2, '094423def9986b8e47457a6592e143f75d67eaa7647c2c346df93d54b983e4e0'),
 'supporting_method_evidence': (4, 'db4e111a865957df036ada510d96050eb5c2b39ed6635c628d48a9a2717ec547'),
 'local_mapping_corrections': (4, '9a1ccc7df1e46b7e8240c3963456fcaf0177d07ad31c051db7d1d5334a1aca1d'),
 'withdrawn_local_conflicts': (1, '8b54b3a107c8f050e3c33cc2125a710a4084b1be287fdb20db8ca392c093ed05')}

BASELINE_ROSTER_UNCHANGED = {'schema_version': '5a4747fcdafec66d8d8812022e3b7d2666d886ec0e0ee7926ec8c50db40f29d5',
 'role_interpretation': '9c1098d16b8ae13b4ef55efb781195831efdc1be68fd601aa8c4d1db41cea7c0'}

BASELINE_QUEUE_PREFIX = (215, 'dcb5ec7e127535410da85c67ac0dfa344430ef4a042b61e7305f3b4d257413ef')

UNCHANGED_APPLICATION_SHA256 = {'site/app.js': '0e420fef54c0dae966d883bb0aeb0fb7af05f79e8803e366ea8ece7e95ed2cd6',
 'site/index.html': '42966eb1409ce029b90d64e77e2c02977971477bade4e2a79414b4a2aef1c571',
 'site/model.js': '087ddd9e9de963f6b4497b3ec50033b1012680057d106a57b490cda4697ec79e',
 'site/rosters.js': '114e47a44f0c4b5cd100eb786604818cacc36ff073f0ccf8e3a20a71cf6e9940',
 'site/styles.css': '6063be4c04726d2dcf01a0442d5dc4d763b99c835a4110bd943799582c14fc34',
 'site/universities.js': '9323dad42fc2c35808ad92888cdffb0ad41edaefaee88c3b57ebb798349c2d62',
 'scripts/build.py': '611d8aefb6e1a043a46426b47b0e35515a87815d1b312604a466c0d210a8e660',
 'scripts/build_rosters.py': 'bc00cde087ec2ba1065b4ac2332fda16918885ec497c2c34dbd319d16f96a7bf',
 'scripts/deterministic_xlsx.py': '65158f54d001eebb2ac586fd820cd69deed8e5c96ed5b117668ced51c11dc300'}

# Reviewed release24 display-only compatibility overlay. Preserve the historical
# pins above; only these four assets may differ after the two fallback labels and
# six shared cache tokens change. All source/data/field pins remain unchanged.
DISPLAY_FIX_APPLICATION_SHA256 = {
    'site/app.js': '5fec893a4b83318d4247299a545f34845d37e076d502bcc9100acec6f73e8f59',
    'site/index.html': '4dde40455944aacd2b3a0a463aee89a64ae2d81ad792a3740dd43c42f005e72f',
    'site/model.js': 'e932f5275311fc833c9ece21a82ade6713987e90e8c4269e199b1a3f07aaed8c',
    'site/rosters.js': 'c3db53692e71b9874be1c003fe6e29a92f4b607a67730d2263c61de62b4874bd',
}

SOURCE_FINGERPRINTS = {'nju-chengxin': (49, 'a5c2a222999e64cf3b50590c05c4a828e7d47555fa84311e72f3378012d9f8c7'),
 'nju-chenpengfei': (68, 'b0cd451e1d5e0e10aa9c34a1be5bb9bbe770646b3c20d5d838a9431db8156ff4'),
 'nju-chenyang': (79, '1eaa226cd2783a6e17ca62779e835f2de12ed8e704a739b8487b5af8e79991f3'),
 'hnu-chenzengping': (91, '671853a6778e14dc49eaae9d3d2b22f5868d4300303c59074b903fa20bd17370'),
 'hnu-chuxia': (95, '0badfe8f8107b8e8f4075bb7de5794c204cd0a03019e4f3b7c10432f638fe475'),
 'hnu-huanshuangyan': (103, '1a232d6284514dafb04aa05871578c559babad2d667998a52c48acfa01993d4d')}

ORG_FINGERPRINTS = {'nju-astronomy-space-science': (104, '2f63c77b68afac6d6d66a0e641f3529844074c17fe187c460c5591957a43062c'),
 'hnu-chemistry': (104, '2d853db2be43f251c572a7778a949025cf62587539460cdbb6de8d96fe54224d'),
 'hnu-chemistry-department': (104, '20095f9d78151246109498d2db03a2b1fad4a6f2fed33a89903194480fd0bb4b')}

EXPECTED_EDUCATION = {'nju-chengxin': [('nju-chengxin-education-1',
                   '学士',
                   '浙江工业大学（CV原文 Zhejiang University of Technology）',
                   '2003-09',
                   '2007-06',
                   'degree_study',
                   None,
                   None),
                  ('nju-chengxin-education-2', '博士', '南京大学', '2007-09', '2012-06', 'degree', True, '2012'),
                  ('nju-chengxin-education-3',
                   None,
                   'George Mason University',
                   '2010-10',
                   '2011-12',
                   'joint_training',
                   False,
                   None),
                  ('nju-chengxin-education-4',
                   None,
                   'Max Planck Institute for Solar System Research',
                   '2019-02',
                   '2022-07',
                   'academic_visit',
                   False,
                   None)],
 'nju-chenpengfei': [('nju-chenpengfei-education-1', '理学博士', '南京大学', None, '1999', 'degree', True, '1999'),
                     ('nju-chenpengfei-education-2',
                      None,
                      '京都大学天文系',
                      '1999-10',
                      '2000-04',
                      'postdoctoral',
                      False,
                      None),
                     ('nju-chenpengfei-education-3',
                      None,
                      '京都大学天文系',
                      '2001-11',
                      '2002-11',
                      'postdoctoral',
                      False,
                      None),
                     ('nju-chenpengfei-education-4',
                      None,
                      '阿拉巴马州立大学（来源原称）',
                      None,
                      None,
                      'academic_visit',
                      False,
                      None),
                     ('nju-chenpengfei-education-5',
                      None,
                      '高山天文台',
                      None,
                      None,
                      'academic_visit',
                      False,
                      None),
                     ('nju-chenpengfei-education-6',
                      None,
                      '圣安德鲁斯大学',
                      None,
                      None,
                      'academic_visit',
                      False,
                      None),
                     ('nju-chenpengfei-education-7',
                      None,
                      '马普研究所（具体研究所未披露）',
                      None,
                      None,
                      'academic_visit',
                      False,
                      None),
                     ('nju-chenpengfei-education-8',
                      None,
                      '新泽西理工学院',
                      '2008-03',
                      '2008-06',
                      'academic_visit',
                      False,
                      None)],
 'nju-chenyang': [('nju-chenyang-education-1', '本科', '南京大学', '1982', '1986', 'degree_study', None, None),
                  ('nju-chenyang-education-2', '理学博士', '南京大学', None, '1993', 'degree', True, '1993'),
                  ('nju-chenyang-education-3',
                   None,
                   'Arcetri天文台',
                   '1991',
                   '1992',
                   'joint_training',
                   False,
                   None),
                  ('nju-chenyang-education-4',
                   None,
                   'Harvard-Smithsonian Center for Astrophysics',
                   '1999',
                   '2000',
                   'academic_visit',
                   False,
                   None),
                  ('nju-chenyang-education-5',
                   None,
                   'University of Massachusetts Amherst',
                   '2003',
                   '2003',
                   'academic_visit',
                   False,
                   None),
                  ('nju-chenyang-education-6',
                   None,
                   'University of Massachusetts Amherst',
                   '2008',
                   '2008',
                   'academic_visit',
                   False,
                   None)],
 'hnu-chenzengping': [('hnu-chenzengping-education-1',
                       '理学学士',
                       '湖南大学',
                       '1991-09',
                       '1995-06',
                       'degree_study',
                       None,
                       None),
                      ('hnu-chenzengping-education-2',
                       '理学硕士',
                       '湖南大学',
                       '1995-09',
                       '1998-06',
                       'degree_study',
                       None,
                       None),
                      ('hnu-chenzengping-education-3',
                       '理学博士',
                       '湖南大学',
                       '1998-09',
                       '2001-05',
                       'degree_study',
                       True,
                       '2001'),
                      ('hnu-chenzengping-education-4',
                       None,
                       '纽卡斯尔大学',
                       '2002-10',
                       '2008-11',
                       'postdoctoral_research',
                       False,
                       None)],
 'hnu-chuxia': [('hnu-chuxia-education-1', '学士', '湖南大学', '1989-09', '1993-06', 'degree_study', None, None),
                ('hnu-chuxia-education-2', '硕士', '湖南大学', '1993-09', '1996-06', 'degree_study', None, None),
                ('hnu-chuxia-education-3', '博士', '湖南大学', '1996-09', '1999-06', 'degree_study', None, None)],
 'hnu-huanshuangyan': [('hnu-huanshuangyan-education-1',
                        None,
                        '中南民族学院',
                        '1994-09',
                        '1998-06',
                        'degree_study',
                        None,
                        None),
                       ('hnu-huanshuangyan-education-2',
                        '硕士',
                        '华中师范大学',
                        '1998-09',
                        '2001-06',
                        'degree_study',
                        None,
                        None),
                       ('hnu-huanshuangyan-education-3',
                        '博士',
                        '湖南大学',
                        '2001-09',
                        '2004-07',
                        'degree_study',
                        None,
                        None),
                       ('hnu-huanshuangyan-education-4',
                        None,
                        'McMaster University',
                        '2012-02',
                        '2013-03',
                        'academic_visit',
                        False,
                        None)]}

EXPECTED_COUNTS = {'nju-chengxin': {'education': 4,
                  'appointments': 4,
                  'mentors': 2,
                  'publications': 3,
                  'pending_education_candidates': 1,
                  'academic_service': 6,
                  'honors': 9,
                  'research_projects': 7,
                  'teaching': 4},
 'nju-chenpengfei': {'education': 8,
                     'appointments': 4,
                     'mentors': 1,
                     'publications': 1,
                     'pending_education_candidates': 0,
                     'academic_service': 15,
                     'honors': 5,
                     'research_projects': 2,
                     'teaching': 3},
 'nju-chenyang': {'education': 6,
                  'appointments': 4,
                  'mentors': 0,
                  'publications': 1,
                  'pending_education_candidates': 1,
                  'academic_service': 8,
                  'honors': 5,
                  'research_projects': 2,
                  'teaching': 3},
 'hnu-chenzengping': {'education': 4,
                      'appointments': 2,
                      'mentors': 0,
                      'publications': 1,
                      'pending_education_candidates': 0,
                      'academic_service': 0,
                      'honors': 0,
                      'research_projects': 0,
                      'teaching': 0},
 'hnu-chuxia': {'education': 3,
                'appointments': 5,
                'mentors': 0,
                'publications': 1,
                'pending_education_candidates': 0,
                'academic_service': 0,
                'honors': 0,
                'research_projects': 0,
                'teaching': 0},
 'hnu-huanshuangyan': {'education': 4,
                       'appointments': 4,
                       'mentors': 0,
                       'publications': 1,
                       'pending_education_candidates': 0,
                       'academic_service': 0,
                       'honors': 0,
                       'research_projects': 0,
                       'teaching': 0}}

EXPECTED_SCOPES = {'nju': '天文与空间科学学院新版官方教学科研人员“教授”栏目29人，按原生页面infolist渲染原序选首3：程鑫、陈鹏飞、陈阳。未按履历、海外经历或师承丰富度筛人，未按内部listorder字段自行重排。仅3人的部分公开职业资料，非全院或全校完成；名录现列不等于全职人事证明。旧版教授页顺序及总数不同，本次统一使用新版栏目。',
 'hnu': '化学化工学院教师名册中“化学系分析化学学科—教授”栏目正文20人，原序前3陈增萍、楚霞、宦双燕；不按学历、海外经历或资料丰富度筛人。教授是栏目类别，不是组织节点。仅3人的部分职业资料，非全学科、全院或全校完成；名录现列不等人事在岗证明。'}

EXPECTED_LIMITATIONS = {'nju': ['天文与空间科学学院教授名录29人，仅原序首3样本；原有6名NJU样本不重复。',
         '原生HTML/JS读取和网页正文读取均为普通公开访问，没有浏览器、远程写入或绕过拒绝。',
         '本包保留事实字段与逐字段直接来源，不发布原始网页/PDF全文、联系方式、出生日期、家庭信息、学生名录或排名。',
         '职称、行政职务、学术服务、科研项目、访问经历分层；博士后/访问/联合培养明确非授位；未明学位授予用null。',
         '程鑫本科院校与联合培养开始月存在来源冲突，冲突保留且本科院校退出教育筛选；不选择性覆盖。',
         '丁明德、Jie Zhang博士指导关系明确且有当前官方导航；方成博士指导明确但当前导航未核；陈阳导师未知。',
         '校区全部未知，不以办公室、页脚地址、学校整体校区或访问地推断。',
         '出版物每人最多3篇；仅书目元数据，部分出版者链接未能核验；不把共同作者、委员会、项目搭档当导师。'],
 'hnu': ['化学化工学院教师名册中“化学系分析化学学科—教授”栏目正文20人，原序前3陈增萍、楚霞、宦双燕；不按学历、海外经历或资料丰富度筛人。教授是栏目类别，不是组织节点。仅3人的部分职业资料，非全学科、全院或全校完成；名录现列不等人事在岗证明。',
         '只含公开职业事实及每人1条精选书目；不含联系方式、出生、政治背景、学生名单或完整网页/API正文。',
         '学位授予采用true/null/false三态：陈增萍博士由官方报告明确获授；其余标签不自动确认为已获学位；博士后和访问为false。',
         '没有本批人物本人具名导师的直接阶段性证据；不从合著、招生资格、项目或第三人的答辩委员名单生成师承。',
         '楚霞实验室副主任采用明确日期历史履历，未将简介未更新表述升为当前任职。',
         '宦双燕本科仅层次，保留degree=null；不由任职或博士经历推断本科学位。',
         '个人页应用壳不等正文；本批正文来自同源公开只读接口HTTP200，未用浏览器、登录或规避访问限制。',
         '校区、全职人事状态及出版者单篇入口未核。']}

EXPECTED_NAMES = {'nju': ['程鑫', '陈鹏飞', '陈阳'], 'hnu': ['陈增萍', '楚霞', '宦双燕']}

EXPECTED_IDS = {'nju': ['nju-chengxin', 'nju-chenpengfei', 'nju-chenyang'],
 'hnu': ['hnu-chenzengping', 'hnu-chuxia', 'hnu-huanshuangyan']}

SOURCE_SHAPES = [
    ['value'],  # 0
    ['array', []],  # 1
    ['object', [['institution', 0], ['title', 0], ['start', 0], ['end', 0]]],  # 2
    ['object', [['institution', 0], ['title', 0], ['start', 0], ['end', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 2], ['field_status', 2], ['current', 0], ['status', 0], ['as_of', 0], ['note', 0], ['employment_type', 0], ['personnel_status_verified', 0]]],  # 3
    ['array', [3, 3, 3, 3]],  # 4
    ['object', [['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['degree_awarded', 0], ['degree_awarded_date', 0]]],  # 5
    ['object', [['event_id', 0], ['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 5], ['field_status', 5], ['country', 0], ['region', 0], ['continent', 0], ['geography_basis', 0], ['geography_status', 0], ['event_kind', 0], ['note', 0], ['include_in_education_filters', 0]]],  # 6
    ['array', [6, 6, 6, 6]],  # 7
    ['object', [['name', 0], ['type', 0], ['education_stage', 0], ['education_event_id', 0], ['start', 0], ['end', 0], ['aliases', 0], ['current_institution', 0], ['official_url', 0], ['navigation_verified', 0]]],  # 8
    ['array', [0, 0, 0]],  # 9
    ['array', [0]],  # 10
    ['object', [['name', 0], ['type', 0], ['education_stage', 0], ['education_event_id', 0], ['start', 0], ['end', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 8], ['field_status', 8], ['evidence_note', 0], ['relationship_status', 0], ['include_in_mentorship_graph', 0], ['navigation_verified', 0], ['navigation_status', 0], ['employment_status', 0], ['current_institution', 0], ['official_url', 0], ['person_id', 0], ['navigation_source_ids', 9], ['note', 0], ['aliases', 10]]],  # 11
    ['object', [['name', 0], ['type', 0], ['education_stage', 0], ['education_event_id', 0], ['start', 0], ['end', 0], ['current_institution', 0], ['official_url', 0], ['navigation_verified', 0]]],  # 12
    ['object', [['name', 0], ['type', 0], ['education_stage', 0], ['education_event_id', 0], ['start', 0], ['end', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 12], ['field_status', 12], ['evidence_note', 0], ['relationship_status', 0], ['include_in_mentorship_graph', 0], ['navigation_verified', 0], ['navigation_status', 0], ['employment_status', 0], ['current_institution', 0], ['official_url', 0], ['person_id', 0], ['navigation_source_ids', 10], ['note', 0]]],  # 13
    ['array', [11, 13]],  # 14
    ['object', [['title', 0], ['year', 0], ['venue', 0], ['volume', 0], ['pages_or_article_number', 0], ['doi', 0], ['type', 0]]],  # 15
    ['object', [['title', 0], ['year', 0], ['venue', 0], ['volume', 0], ['pages_or_article_number', 0], ['doi', 0], ['type', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 15], ['field_status', 15], ['url', 0], ['identity_evidence', 0], ['link_status', 0], ['individual_article_link_verified', 0], ['note', 0]]],  # 16
    ['object', [['title', 0], ['year', 0], ['venue', 0], ['volume', 0], ['pages_or_article_number', 0], ['doi', 0], ['type', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 15], ['field_status', 15], ['url', 0], ['identity_evidence', 0], ['link_status', 0], ['individual_article_link_verified', 0], ['note', 0], ['identity_source_ids', 9]]],  # 17
    ['array', [16, 17, 16]],  # 18
    ['object', [['name', 0]]],  # 19
    ['object', [['name', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 19], ['field_status', 19]]],  # 20
    ['array', [20, 20, 20]],  # 21
    ['object', [['organization', 0], ['role', 0], ['start', 0], ['end', 0], ['status', 0], ['as_of', 0]]],  # 22
    ['object', [['organization', 0], ['role', 0], ['start', 0], ['end', 0], ['status', 0], ['as_of', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 22], ['field_status', 22]]],  # 23
    ['object', [['organization', 0], ['role', 0], ['start', 0], ['end', 0], ['status', 0]]],  # 24
    ['object', [['organization', 0], ['role', 0], ['start', 0], ['end', 0], ['status', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 24], ['field_status', 24]]],  # 25
    ['array', [23, 23, 25, 25, 25, 25]],  # 26
    ['object', [['course', 0], ['level', 0], ['period', 0]]],  # 27
    ['object', [['course', 0], ['level', 0], ['period', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 27], ['field_status', 27]]],  # 28
    ['array', [28, 28, 28, 28]],  # 29
    ['object', [['name', 0], ['year', 0]]],  # 30
    ['object', [['name', 0], ['year', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 30], ['field_status', 30]]],  # 31
    ['object', [['name', 0], ['year', 0], ['note', 0]]],  # 32
    ['object', [['name', 0], ['year', 0], ['note', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 32], ['field_status', 32]]],  # 33
    ['array', [31, 31, 31, 31, 33, 33, 33, 33, 33]],  # 34
    ['object', [['name', 0], ['year', 0], ['role', 0], ['start', 0], ['end', 0]]],  # 35
    ['object', [['name', 0], ['year', 0], ['role', 0], ['start', 0], ['end', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 35], ['field_status', 35]]],  # 36
    ['object', [['name', 0], ['role', 0], ['start', 0], ['end', 0], ['note', 0]]],  # 37
    ['object', [['name', 0], ['role', 0], ['start', 0], ['end', 0], ['note', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 37], ['field_status', 37]]],  # 38
    ['array', [36, 36, 36, 38, 38, 38, 38]],  # 39
    ['object', [['id', 0], ['url', 0], ['evidence_status', 0], ['retrieved_at', 0], ['source_type', 0], ['access_note', 0]]],  # 40
    ['object', [['id', 0], ['url', 0], ['evidence_status', 0], ['retrieved_at', 0], ['source_type', 0]]],  # 41
    ['array', [40, 41, 40, 40, 41, 40, 40, 40, 40, 40]],  # 42
    ['array', [0, 0, 0, 0, 0]],  # 43
    ['array', [0, 0]],  # 44
    ['object', [['status', 0], ['source_url', 0], ['profile_url', 0], ['source_ids', 44]]],  # 45
    ['object', [['name', 0], ['school_id', 0], ['department_id', 0], ['organization_id', 0], ['listed_role', 0]]],  # 46
    ['object', [['candidate_type', 0], ['institution', 0], ['degree', 0], ['field', 0], ['end', 0], ['degree_awarded', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['verification_status', 0], ['include_in_education_filters', 0], ['note', 0]]],  # 47
    ['array', [47]],  # 48
    ['object', [['person_id', 0], ['name', 0], ['school_id', 0], ['department_id', 0], ['organization_id', 0], ['campus_ids', 1], ['campus', 0], ['department', 0], ['roster_entry_id', 0], ['official_url', 0], ['listed_role', 0], ['listed_role_source_url', 0], ['listed_role_source_id', 0], ['listed_role_basis', 0], ['appointments', 4], ['education', 7], ['mentors', 14], ['committee_members', 1], ['publications', 18], ['research_fields', 21], ['academic_service', 26], ['teaching', 29], ['honors', 34], ['research_projects', 39], ['sources', 42], ['unknowns', 43], ['identity_evidence', 45], ['professional_field_provenance', 46], ['profile_status', 0], ['pending_education_candidates', 48], ['pending_mentor_candidates', 1], ['pending_appointment_candidates', 1], ['pending_professional_fact_candidates', 1]]],  # 49
    ['object', [['start', 0], ['end', 0], ['institutions', 0], ['scope', 0], ['individual_institution_period_known', 0]]],  # 50
    ['object', [['start', 0], ['end', 0], ['institutions', 44], ['scope', 0], ['individual_institution_period_known', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 50], ['field_status', 50]]],  # 51
    ['object', [['event_id', 0], ['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 5], ['field_status', 5], ['country', 0], ['region', 0], ['continent', 0], ['geography_basis', 0], ['geography_status', 0], ['event_kind', 0], ['note', 0], ['include_in_education_filters', 0], ['source_combined_period', 51]]],  # 52
    ['array', [6, 6, 6, 52, 52, 52, 52, 6]],  # 53
    ['object', [['name', 0], ['type', 0], ['education_stage', 0], ['education_event_id', 0], ['start', 0], ['end', 0]]],  # 54
    ['object', [['name', 0], ['type', 0], ['education_stage', 0], ['education_event_id', 0], ['start', 0], ['end', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 54], ['field_status', 54], ['evidence_note', 0], ['relationship_status', 0], ['include_in_mentorship_graph', 0], ['navigation_verified', 0], ['navigation_status', 0], ['employment_status', 0], ['current_institution', 0], ['official_url', 0], ['person_id', 0], ['navigation_source_ids', 1], ['note', 0]]],  # 55
    ['array', [55]],  # 56
    ['array', [0, 0, 0, 0]],  # 57
    ['object', [['title', 0], ['year', 0], ['venue', 0], ['volume', 0], ['pages_or_article_number', 0], ['doi', 0], ['type', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 15], ['field_status', 15], ['url', 0], ['identity_evidence', 0], ['link_status', 0], ['individual_article_link_verified', 0], ['note', 0], ['identity_source_ids', 57]]],  # 58
    ['array', [58]],  # 59
    ['array', [20, 20, 20, 20, 20, 20]],  # 60
    ['object', [['organization', 0], ['role', 0], ['start', 0], ['end', 0], ['status', 0], ['as_of', 0], ['note', 0]]],  # 61
    ['object', [['organization', 0], ['role', 0], ['start', 0], ['end', 0], ['status', 0], ['as_of', 0], ['note', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 61], ['field_status', 61]]],  # 62
    ['array', [25, 23, 23, 23, 23, 23, 23, 62, 62, 62, 62, 62, 62, 62, 62]],  # 63
    ['array', [28, 28, 28]],  # 64
    ['array', [31, 31, 31, 31, 31]],  # 65
    ['array', [36, 36]],  # 66
    ['array', [40, 41, 40, 41, 40, 40, 40]],  # 67
    ['object', [['person_id', 0], ['name', 0], ['school_id', 0], ['department_id', 0], ['organization_id', 0], ['campus_ids', 1], ['campus', 0], ['department', 0], ['roster_entry_id', 0], ['official_url', 0], ['listed_role', 0], ['listed_role_source_url', 0], ['listed_role_source_id', 0], ['listed_role_basis', 0], ['appointments', 4], ['education', 53], ['mentors', 56], ['committee_members', 1], ['publications', 59], ['research_fields', 60], ['academic_service', 63], ['teaching', 64], ['honors', 65], ['research_projects', 66], ['sources', 67], ['unknowns', 43], ['identity_evidence', 45], ['professional_field_provenance', 46], ['profile_status', 0], ['pending_education_candidates', 1], ['pending_mentor_candidates', 1], ['pending_appointment_candidates', 1], ['pending_professional_fact_candidates', 1]]],  # 68
    ['array', [6, 6, 6, 6, 6, 6]],  # 69
    ['array', [16]],  # 70
    ['array', [20, 20, 20, 20]],  # 71
    ['array', [25, 25, 25, 25, 25, 25, 25, 25]],  # 72
    ['object', [['name', 0], ['year', 0], ['role', 0], ['start', 0], ['end', 0], ['note', 0]]],  # 73
    ['object', [['name', 0], ['year', 0], ['role', 0], ['start', 0], ['end', 0], ['note', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['field_provenance', 73], ['field_status', 73]]],  # 74
    ['array', [36, 74]],  # 75
    ['array', [40, 41, 40, 40, 40]],  # 76
    ['object', [['candidate_type', 0], ['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['degree_awarded', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['verification_status', 0], ['include_in_education_filters', 0], ['note', 0]]],  # 77
    ['array', [77]],  # 78
    ['object', [['person_id', 0], ['name', 0], ['school_id', 0], ['department_id', 0], ['organization_id', 0], ['campus_ids', 1], ['campus', 0], ['department', 0], ['roster_entry_id', 0], ['official_url', 0], ['listed_role', 0], ['listed_role_source_url', 0], ['listed_role_source_id', 0], ['listed_role_basis', 0], ['appointments', 4], ['education', 69], ['mentors', 1], ['committee_members', 1], ['publications', 70], ['research_fields', 71], ['academic_service', 72], ['teaching', 64], ['honors', 65], ['research_projects', 75], ['sources', 76], ['unknowns', 57], ['identity_evidence', 45], ['professional_field_provenance', 46], ['profile_status', 0], ['pending_education_candidates', 78], ['pending_mentor_candidates', 1], ['pending_appointment_candidates', 1], ['pending_professional_fact_candidates', 1]]],  # 79
    ['object', [['institution', 0], ['title', 0], ['start', 0], ['end', 0], ['current', 0], ['status', 0], ['as_of', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['note', 0], ['employment_type', 0], ['personnel_status_verified', 0], ['field_provenance', 2], ['field_status', 2]]],  # 80
    ['array', [80, 80]],  # 81
    ['object', [['event_id', 0], ['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['country', 0], ['region', 0], ['continent', 0], ['geography_basis', 0], ['geography_status', 0], ['event_kind', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['note', 0], ['include_in_education_filters', 0], ['field_provenance', 5], ['field_status', 5]]],  # 82
    ['object', [['event_id', 0], ['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['country', 0], ['region', 0], ['continent', 0], ['geography_basis', 0], ['geography_status', 0], ['event_kind', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['note', 0], ['include_in_education_filters', 0], ['field_provenance', 5], ['field_status', 5], ['supporting_source_ids', 10]]],  # 83
    ['array', [82, 82, 83, 82]],  # 84
    ['object', [['title', 0], ['year', 0], ['venue', 0], ['volume', 0], ['pages_or_article_number', 0], ['doi', 0], ['url', 0], ['type', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['identity_evidence', 0], ['link_status', 0], ['individual_article_link_verified', 0], ['note', 0], ['field_provenance', 15]]],  # 85
    ['array', [85]],  # 86
    ['array', [41, 40, 40]],  # 87
    ['object', [['status', 0], ['source_url', 0], ['profile_url', 0], ['source_ids', 44], ['note', 0]]],  # 88
    ['object', [['name', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0]]],  # 89
    ['array', [89, 89, 89, 89, 89]],  # 90
    ['object', [['person_id', 0], ['name', 0], ['school_id', 0], ['department_id', 0], ['organization_id', 0], ['campus_ids', 1], ['campus', 0], ['department', 0], ['roster_entry_id', 0], ['official_url', 0], ['listed_role', 0], ['listed_role_source_url', 0], ['listed_role_source_id', 0], ['listed_role_basis', 0], ['appointments', 81], ['education', 84], ['mentors', 1], ['committee_members', 1], ['publications', 86], ['sources', 87], ['unknowns', 43], ['identity_evidence', 88], ['professional_field_provenance', 46], ['profile_status', 0], ['pending_education_candidates', 1], ['pending_mentor_candidates', 1], ['pending_appointment_candidates', 1], ['pending_professional_fact_candidates', 1], ['research_fields', 90]]],  # 91
    ['array', [80, 80, 80, 80, 80]],  # 92
    ['array', [82, 82, 82]],  # 93
    ['array', [89, 89, 89]],  # 94
    ['object', [['person_id', 0], ['name', 0], ['school_id', 0], ['department_id', 0], ['organization_id', 0], ['campus_ids', 1], ['campus', 0], ['department', 0], ['roster_entry_id', 0], ['official_url', 0], ['listed_role', 0], ['listed_role_source_url', 0], ['listed_role_source_id', 0], ['listed_role_basis', 0], ['appointments', 92], ['education', 93], ['mentors', 1], ['committee_members', 1], ['publications', 86], ['sources', 87], ['unknowns', 57], ['identity_evidence', 88], ['professional_field_provenance', 46], ['profile_status', 0], ['pending_education_candidates', 1], ['pending_mentor_candidates', 1], ['pending_appointment_candidates', 1], ['pending_professional_fact_candidates', 1], ['research_fields', 94]]],  # 95
    ['array', [80, 80, 80, 80]],  # 96
    ['object', [['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['education_level', 0]]],  # 97
    ['object', [['event_id', 0], ['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['country', 0], ['region', 0], ['continent', 0], ['geography_basis', 0], ['geography_status', 0], ['event_kind', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['note', 0], ['include_in_education_filters', 0], ['field_provenance', 97], ['field_status', 97], ['education_level', 0]]],  # 98
    ['object', [['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['role', 0]]],  # 99
    ['object', [['event_id', 0], ['degree', 0], ['institution', 0], ['start', 0], ['end', 0], ['field', 0], ['country', 0], ['region', 0], ['continent', 0], ['geography_basis', 0], ['geography_status', 0], ['event_kind', 0], ['degree_awarded', 0], ['degree_awarded_date', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0], ['note', 0], ['include_in_education_filters', 0], ['field_provenance', 99], ['field_status', 99], ['role', 0]]],  # 100
    ['array', [98, 82, 82, 100]],  # 101
    ['array', [41, 40]],  # 102
    ['object', [['person_id', 0], ['name', 0], ['school_id', 0], ['department_id', 0], ['organization_id', 0], ['campus_ids', 1], ['campus', 0], ['department', 0], ['roster_entry_id', 0], ['official_url', 0], ['listed_role', 0], ['listed_role_source_url', 0], ['listed_role_source_id', 0], ['listed_role_basis', 0], ['appointments', 96], ['education', 101], ['mentors', 1], ['committee_members', 1], ['publications', 86], ['sources', 102], ['unknowns', 57], ['identity_evidence', 88], ['professional_field_provenance', 46], ['profile_status', 0], ['pending_education_candidates', 1], ['pending_mentor_candidates', 1], ['pending_appointment_candidates', 1], ['pending_professional_fact_candidates', 1], ['research_fields', 94]]],  # 103
    ['object', [['id', 0], ['name', 0], ['type', 0], ['school_id', 0], ['parent_id', 0], ['source_url', 0], ['source_id', 0], ['evidence_status', 0]]],  # 104
]

NODE_CONSUMERS = "\nimport fs from 'node:fs';\nimport {pathToFileURL} from 'node:url';\nconst input=JSON.parse(fs.readFileSync(0,'utf8')), root=input.root;\nconst model=await import(pathToFileURL(root+'/site/model.js'));\nconst roster=await import(pathToFileURL(root+'/site/rosters.js'));\nconst db=input.atlas, r=input.rosters, before=JSON.stringify([db,r]);\nconst people=new Map([...db.people,...db.pending_people].map(p=>[p.person_id,p]));\nconst scoped={...db,rosters:r}, out={people:{},rosters:{},schoolStats:{},schoolPages:{}};\nfor(const id of input.ids){\n const p=people.get(id),e=r.entries.find(e=>e.id===p.roster_entry_id),main={innerHTML:''};\n if(!p||!e)throw Error('Missing release record '+id);\n roster.renderRosters(scoped,main,e.id,p.school_id);\n out.rosters[id]={matched:roster.rosterIdentityMatched(scoped,e),status:roster.rosterStatusLabel(scoped,e),html:main.innerHTML};\n out.people[id]={\n  stages:Object.fromEntries(['bachelor','master','doctor'].map(s=>[s,model.stages(p,s).map(e=>e.event_id)])),\n  ug:Object.fromEntries(['yes','no','unknown','not_recorded'].map(v=>[v,model.matches(p,{ug:v},db.org_units)])),\n  doctorate:Object.fromEntries(['mainland','hmt','overseas','unknown','not_recorded'].map(v=>[v,model.matches(p,{doctor:v},db.org_units)])),\n  officialUrl:model.officialProfileUrl(p),\n  mentors:p.mentors.map(m=>({name:m.name,url:model.mentorOfficialUrl(m),route:model.mentorProfileRoute(m,db.people),events:model.mentorEvents(p,m).map(e=>e.event_id)})),\n  candidates:(p.pending_mentor_candidates||[]).map(m=>({name:m.name,url:model.mentorOfficialUrl(m),route:model.mentorProfileRoute(m,db.people)})),\n  labels:p.education.map(model.educationStageLabel)\n };\n}\nfor(const school of input.schools)out.schoolStats[school]=model.schoolStats(scoped,school);\nconst appAtlas=structuredClone(db),appRosters=structuredClone(r);\nappAtlas.build_id=appRosters.build_id='rotation002-wave2-in-memory-test';\nappAtlas.exports=input.exports;\nout.directories={};\nconst nodes=new Map(),listeners={};\nconst node=key=>{\n if(!nodes.has(key))nodes.set(key,{innerHTML:'',textContent:'',value:'',style:{},dataset:{height:'320'},clientWidth:960,addEventListener(){},click(){this.onclick?.()},querySelector(selector){return node(selector)},scrollIntoView(){}});\n return nodes.get(key);\n};\nglobalThis.document={querySelector:node,querySelectorAll:()=>[],getElementById:id=>node('#'+id)};\nglobalThis.location={hash:'#/',href:'https://atlas.test/site/index.html'};\nglobalThis.window={scrollTo(){},addEventListener(name,handler){listeners[name]=handler}};\nconst {createAtlasFetch}=await import(pathToFileURL(root+'/tests/atlas-test-io.js'));\nglobalThis.fetch=createAtlasFetch(appAtlas,appRosters);\nawait import(pathToFileURL(root+'/site/app.js'));\nif(!listeners.hashchange)throw Error('Actual app failed to initialize: '+node('#main').innerHTML);\nconst navigate=hash=>{for(const n of nodes.values())n.innerHTML='';location.hash=hash;listeners.hashchange()};\nfor(const id of input.ids){\n const p=people.get(id),row=out.people[id];\n navigate('#/person/'+encodeURIComponent(id));\n row.profileMain=node('#main').innerHTML;row.profileContent=node('#profile-content').innerHTML;\n if(p.detail_count_eligible){\n  navigate('#/person/'+encodeURIComponent(id)+'/lineage');\n  row.lineage=node('#profile-content').innerHTML;row.graph=node('#mind-stage').innerHTML;\n }\n}\nfor(const school of input.schools){\n navigate('#/school/'+school);out.schoolPages[school]=node('#main').innerHTML;\n navigate('#/people?school='+school);out.directories[school]=node('#results').innerHTML;\n}\nout.unchanged=JSON.stringify([db,r])===before;\nprocess.stdout.write(JSON.stringify(out));\n"


def project_source_shape(value, node_id):
    """Refuse lost keys, changed array lengths, or changed original leaf values."""
    node = SOURCE_SHAPES[node_id]
    if node[0] == 'value':
        return value
    if node[0] == 'object':
        assert isinstance(value, dict), 'Reviewed object changed type'
        return {key: project_source_shape(value[key], child) for key, child in node[1]}
    assert isinstance(value, list) and len(value) == len(node[1]), 'Reviewed collection changed length/type'
    return [project_source_shape(item, child) for item, child in zip(value, node[1])]


def source_view(person):
    """Reverse only explicitly retained source aliases, never infer missing facts."""
    person = copy.deepcopy(person)
    person['profile_status'] = person['source_profile_status']
    for event in person['education']:
        for key in ('region', 'institution_985', 'event_kind'):
            if 'source_' + key in event:
                event[key] = event['source_' + key]
    for mentor in person['mentors']:
        for key in ('employment_status', 'official_url'):
            if 'source_' + key in mentor:
                mentor[key] = mentor['source_' + key]
    return person


def assert_source_closure(person):
    sources = {row['id']: row['url'] for row in person['sources']}
    assert len(sources) == len(person['sources']), 'Duplicate source ID'
    events = {event['event_id'] for event in person['education']}
    # These two reviewed candidates intentionally have NO event_id.
    candidates = {event['event_id'] for event in person.get('pending_education_candidates', []) if event.get('event_id')}
    assert not events & candidates, 'Candidate event promoted to main'

    def check(value, path='person'):
        if isinstance(value, dict):
            for key, child in value.items():
                if (key == 'source_id' or key.endswith('_source_id')) and child is not None:
                    assert child in sources, (path, key, 'unknown source', child)
                    url_key = key[:-3] + '_url'
                    if value.get(url_key) is not None:
                        assert value[url_key] == sources[child], (path, key, 'source URL mismatch')
                if (key == 'source_ids' or key.endswith('_source_ids')) and child is not None:
                    assert isinstance(child, list) and all(item in sources for item in child), (path, key)
                if key in ('field_provenance', 'professional_field_provenance'):
                    for ref in child.values():
                        refs = ref if isinstance(ref, list) else [ref]
                        assert all(item is None or isinstance(item, str) and item in sources for item in refs), (path, key)
                else:
                    check(child, path + '/' + key)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                check(child, path + '/' + str(index))
    check(person)
    for mentor in person['mentors']:
        assert mentor.get('education_event_id') in events, 'Main mentor must bind a main education event'
    for candidate in person.get('pending_education_candidates', []):
        assert candidate.get('include_in_education_filters') is False, 'Candidate filter promotion'
    return True


class Rotation002Wave2ReleaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = release25_view(load_atlas(ROOT / 'data/atlas.json'), 'data/atlas.json')
        cls.r = release25_view(json.loads((ROOT / 'data/rosters.json').read_text()), 'data/rosters.json')
        cls.q = release25_view(json.loads((ROOT / 'progress/work-queue.json').read_text()), 'progress/work-queue.json')
        cls.ps = {p['person_id']: p for p in cls.d['people'] + cls.d['pending_people']}
        cls.new = [cls.ps[pid] for pid in SOURCE_FINGERPRINTS]
        cls.entries = {e['source_person_id']: e for e in cls.r['entries'] if e.get('source_person_id') in SOURCE_FINGERPRINTS}
        cls.js = None

    @classmethod
    def consumers(cls, atlas=None):
        if atlas is None and cls.js is not None:
            return cls.js
        result = subprocess.run(
            ['node', '--input-type=module', '-e', NODE_CONSUMERS],
            input=json.dumps({'root': str(ROOT), 'atlas': atlas or cls.d, 'rosters': cls.r,
                              'ids': list(SOURCE_FINGERPRINTS), 'schools': list(EXPECTED_IDS),
                              'exports': build.department_exports(cls.d)}),
            text=True, capture_output=True, timeout=45, cwd=ROOT,
        )
        if result.returncode:
            raise AssertionError('Actual JS consumers failed:\n' + result.stderr)
        output = json.loads(result.stdout)
        if atlas is None:
            cls.js = output
        return output

    def tables(self, pid, person=None):
        person = person or self.ps[pid]
        before = copy.deepcopy(person)
        with mock.patch.object(build, 'Workbook', side_effect=AssertionError('Workbook writes prohibited')):
            tables = dict(build.department_tables([person], self.d['updated_at'], self.d['institutions']))
        self.assertEqual(person, before)
        return tables

    def event(self, pid, number):
        return next(e for e in self.ps[pid]['education'] if e['event_id'] == f'{pid}-education-{number}')

    def test_01_original_public_leaf_fingerprints_and_all_provenance_survive(self):
        for pid, (shape, digest) in SOURCE_FINGERPRINTS.items():
            with self.subTest(person=pid):
                self.assertEqual(object_sha(project_source_shape(source_view(self.ps[pid]), shape)), digest)
        orgs = {o['id']: o for o in self.d['org_units']}
        for oid, (shape, digest) in ORG_FINGERPRINTS.items():
            self.assertEqual(object_sha(project_source_shape(orgs[oid], shape)), digest)

    def test_02_all_release22_atlas_records_are_exact_unchanged_prefixes(self):
        for key, (count, digest) in BASELINE_ATLAS_PREFIXES.items():
            self.assertEqual(object_sha(self.d[key][:count]), digest, key)
        for key, digest in BASELINE_ATLAS_UNCHANGED.items():
            self.assertEqual(object_sha(self.d[key]), digest, key)

    def test_03_all_release22_rosters_observations_conflicts_exclusions_are_preserved(self):
        for key, (count, digest) in BASELINE_ROSTER_PREFIXES.items():
            self.assertEqual(object_sha(self.r[key][:count]), digest, key)
        for key, digest in BASELINE_ROSTER_UNCHANGED.items():
            self.assertEqual(object_sha(self.r[key]), digest, key)

    def test_04_queue_original_215_tasks_unchanged_only_two_partial_additions(self):
        count, digest = BASELINE_QUEUE_PREFIX
        self.assertEqual(count, 215)
        self.assertEqual(object_sha(self.q['tasks'][:count]), digest)
        self.assertEqual(len(self.q['tasks']), 222)
        self.assertEqual([t['institution_id'] for t in self.q['tasks'][count:217]], ['nju', 'hnu'])
        for task in self.q['tasks'][count:217]:
            self.assertEqual(task['status'], 'partial')
            self.assertIs(task['scope']['university_complete'], False)
            self.assertEqual(task['campus_ids'], [])

    def test_05_exact_approved_scope_totals_and_no_hust_additions(self):
        self.assertEqual((len(self.d['people']), len(self.d['pending_people']), len(self.d['org_units']), len(self.d['institutions'])), (555,84,229,39))
        self.assertEqual((len(self.r['entries']), len(self.r['source_observations'])), (1721,1759))
        self.assertEqual([p['person_id'] for p in self.d['people'][535:541]], list(SOURCE_FINGERPRINTS))
        self.assertFalse({p['person_id'] for p in self.d['pending_people']} & set(SOURCE_FINGERPRINTS))
        self.assertEqual({p['school_id'] for p in self.d['people'][535:541]}, {'nju','hnu'})
        self.assertEqual({e['school_id'] for e in self.r['entries'][1700:1706]}, {'nju','hnu'})
        self.assertEqual({o['id'] for o in self.d['org_units'][218:221]}, set(ORG_FINGERPRINTS))
        self.assertEqual(len(build.department_exports(self.d)), 111)
        self.assertEqual([b['batch_id'] for b in self.d['profile_batches'][130:132]], ['faculty-batch-nju-rotation-002','faculty-batch-hnu-rotation-002'])
        for p in self.new:
            self.assertTrue(p['detail_count_eligible'])
            self.assertFalse(p['complete_profile_count_eligible'])
            self.assertFalse(p['fields_complete'])
            self.assertEqual(p['profile_status'], 'partial_professional_source_checked')

    def test_06_all_original_main_and_independent_collection_counts(self):
        for p in self.new:
            self.assertEqual({k:len(p.get(k, [])) for k in EXPECTED_COUNTS[p['person_id']]}, EXPECTED_COUNTS[p['person_id']])
        totals = {k:sum(len(p.get(k, [])) for p in self.new) for k in EXPECTED_COUNTS[self.new[0]['person_id']]}
        self.assertEqual(totals, {'education':29,'appointments':23,'mentors':3,'publications':8,'pending_education_candidates':2,'academic_service':29,'honors':19,'research_projects':11,'teaching':10})
        self.assertEqual(Counter(e['degree_awarded'] for p in self.new for e in p['education']), {True:4,None:10,False:15})
        for p in self.new:
            for key in ('professional_research_advisors','project_guidance','research_collaborators','committee_members'):
                self.assertEqual(p.get(key, []), [], (p['person_id'],key))

    def test_07_every_source_chronology_and_normalized_kind_remains_precise(self):
        fields = ('event_id','degree','institution','start','end','event_kind','degree_awarded','degree_awarded_date')
        for p in self.new:
            original = source_view(p)
            self.assertEqual([tuple(e.get(k) for k in fields) for e in original['education']], EXPECTED_EDUCATION[p['person_id']])
            for e, raw in zip(p['education'], original['education']):
                self.assertEqual(e['event_kind'], 'degree' if raw['degree_awarded'] is True else raw['event_kind'])
                self.assertEqual(e['education_event_id'], e['event_id'])
                if e['degree_awarded'] is True:
                    self.assertRegex(e['degree_awarded_date'], r'^\d{4}$')
                    self.assertEqual(e['degree_award_source_id'], e['field_provenance']['degree_awarded'])
                    self.assertEqual(e['degree_award_source_url'], next(s['url'] for s in p['sources'] if s['id']==e['degree_award_source_id']))
                else:
                    self.assertIsNone(e['degree_awarded_date'])

    def test_08_actual_filters_admit_only_four_explicit_doctorates(self):
        js = self.consumers()['people']
        doctors = {'nju-chengxin':2,'nju-chenpengfei':1,'nju-chenyang':2,'hnu-chenzengping':3}
        for p in self.new:
            pid = p['person_id']
            self.assertEqual(js[pid]['stages'], {'bachelor':[],'master':[],'doctor':[f'{pid}-education-{doctors[pid]}'] if pid in doctors else []})
            self.assertEqual(js[pid]['ug'], {'yes':False,'no':False,'unknown':False,'not_recorded':True})
            self.assertEqual(js[pid]['doctorate'], {'mainland':pid in doctors,'hmt':False,'overseas':False,'unknown':False,'not_recorded':pid not in doctors})

    def test_09_idless_candidates_are_kept_outside_main_and_award_is_not_transferred(self):
        candidates = [c for p in self.new for c in p.get('pending_education_candidates', [])]
        self.assertEqual(len(candidates), 2)
        for candidate in candidates:
            self.assertNotIn('event_id', candidate)
            self.assertFalse(candidate['include_in_education_filters'])
        person = self.ps['nju-chengxin']; main = person['education'][0]; candidate = person['pending_education_candidates'][0]
        self.assertEqual(candidate['institution'], 'Zhejiang University of Science and Technology')
        self.assertIs(candidate['degree_awarded'], True)
        self.assertEqual(main['institution'], '浙江工业大学（CV原文 Zhejiang University of Technology）')
        self.assertIsNone(main['degree_awarded'])
        self.assertFalse(main['include_in_education_filters'])
        self.assertNotIn(candidate['institution'], [e['institution'] for e in person['education']])
        self.assertNotIn(candidate['institution'], [row[3] for row in self.tables(person['person_id'])['经历'][1:]])
        self.assertIn(candidate['institution'], self.tables(person['person_id'])['经历'][1][14])
        cy = self.ps['nju-chenyang']['pending_education_candidates'][0]
        self.assertEqual(cy['degree'], '硕士')
        for key in ('institution','start','end','degree_awarded'):
            self.assertIsNone(cy[key])

    def test_10_joint_training_conflict_and_source_precision_are_not_repaired(self):
        p = self.ps['nju-chengxin']; event = self.event(p['person_id'],3)
        self.assertEqual((event['start'],event['end']), ('2010-10','2011-12'))
        self.assertEqual(event['event_kind'], 'joint_training')
        self.assertFalse(event['degree_awarded'])
        self.assertIn('9月', event['note'])
        self.assertIn('10月', ' '.join(p['unknowns']))
        self.assertIn('2010-10', json.dumps(self.tables(p['person_id']), ensure_ascii=False))
        self.assertIsNone(self.event('nju-chenyang',2)['start'])
        self.assertEqual(self.event('nju-chenyang',1)['start'], '1982')

    def test_11_four_combined_host_visits_keep_unknown_individual_dates(self):
        p = self.ps['nju-chenpengfei']
        for number, period in [(4,('2001-07','2001-09')),(5,('2001-07','2001-09')),(6,('2005-11','2006-11')),(7,('2005-11','2006-11'))]:
            event = self.event(p['person_id'],number)
            self.assertIsNone(event['start']); self.assertIsNone(event['end'])
            self.assertFalse(event['degree_awarded'])
            combined = event['source_combined_period']
            self.assertEqual((combined['start'],combined['end']), period)
            self.assertEqual(combined['scope'], 'combined_multi_institution_visit_only')
            self.assertIs(combined['individual_institution_period_known'], False)
            self.assertIn(event['institution'], combined['institutions'])
            row = self.tables(p['person_id'])['经历'][number]
            self.assertEqual(row[3],event['institution'])
            self.assertIsNone(row[4]); self.assertIsNone(row[5])
        self.assertEqual(self.event(p['person_id'],4)['institution'], '阿拉巴马州立大学（来源原称）')
        self.assertEqual(self.event(p['person_id'],7)['institution'], '马普研究所（具体研究所未披露）')

    def test_12_hnu_award_evidence_is_separate_from_study_end_month(self):
        e = self.event('hnu-chenzengping',3)
        self.assertEqual(e['end'], '2001-05')
        self.assertEqual(e['degree_awarded_date'], '2001')
        self.assertEqual(e['source_event_kind'], 'degree_study')
        self.assertEqual(e['event_kind'], 'degree')
        self.assertNotEqual(e['degree_award_source_id'], e['source_id'])
        self.assertNotEqual(e['degree_award_source_url'], e['source_url'])
        row = self.tables('hnu-chenzengping')['经历'][3]
        self.assertEqual((row[11],row[12],row[13]), (True,'2001',e['degree_award_source_url']))
        html = self.consumers()['people']['hnu-chenzengping']['profileContent']
        self.assertIn('授位日期：2001 ', html)
        self.assertIn(e['degree_award_source_url'].replace('&','&amp;'), html)

    def test_13_hnu_historical_director_role_and_undergraduate_level_stay_honest(self):
        p = self.ps['hnu-chuxia']; role = next(a for a in p['appointments'] if a['title']=='副主任')
        self.assertEqual((role['start'],role['end']), ('2017-01','2021-12'))
        self.assertFalse(role['current'])
        self.assertEqual(role['status'], 'historical_source_reported')
        self.assertIsNone(role['as_of'])
        self.assertIn('2021-12',self.consumers()['people'][p['person_id']]['profileContent'])
        e = self.event('hnu-huanshuangyan',1)
        self.assertIsNone(e['degree']); self.assertIsNone(e['degree_awarded'])
        self.assertEqual(e['education_level'], '本科')
        self.assertEqual(e['institution'], '中南民族学院')
        self.assertIsNone(self.tables('hnu-huanshuangyan')['经历'][1][2])
        for p in self.new:
            if p['school_id']=='hnu': self.assertEqual(p['mentors'], [])
            for a in p['appointments']:
                self.assertIs(a.get('personnel_status_verified'), False)

    def test_14_two_verified_current_mentor_shortcuts_preserve_raw_status(self):
        js = self.consumers()['people']; p = self.ps['nju-chengxin']
        for m, consumer in zip(p['mentors'], js[p['person_id']]['mentors']):
            self.assertEqual(m['source_employment_status'], 'official_currently_listed_not_personnel_verified')
            self.assertEqual(m['employment_status'], 'active_officially_listed')
            self.assertTrue(m['navigation_verified']); self.assertTrue(m['quick_link_eligible'])
            self.assertEqual(consumer['url'], m['official_url'])
            self.assertEqual(build.mentor_official_url(m),m['official_url'])
            self.assertEqual(consumer['events'], ['nju-chengxin-education-2'])
            self.assertIsNone(m['start']); self.assertIsNone(m['end'])
            self.assertIn(m['name'], js[p['person_id']]['graph'])
            self.assertIn('不是独立', m['employment_status_note'])
        fang = self.ps['nju-chenpengfei']['mentors'][0]
        self.assertEqual(fang['name'],'方成'); self.assertEqual(fang['employment_status'],'unknown')
        self.assertFalse(fang['navigation_verified']); self.assertFalse(fang['quick_link_eligible'])
        self.assertEqual(js['nju-chenpengfei']['mentors'][0]['url'],'')
        self.assertEqual(build.mentor_official_url(fang),'')
        self.assertIn('方成',js['nju-chenpengfei']['graph'])
        self.assertEqual(self.tables('nju-chenpengfei')['导师'][1][7],'')

    def test_15_recursively_closed_sources_candidates_provenance_and_main_edges(self):
        for p in self.new:
            self.assertTrue(assert_source_closure(p))
            for m in p['mentors']:
                self.assertEqual(m['relationship_kind'],'doctoral_advisor')
                self.assertTrue(m['include_in_mentorship_graph'])
                self.assertTrue(m['include_in_verified_mentor_count'])

    def test_16_teaching_source_levels_survive_without_becoming_degrees(self):
        cx = self.ps['nju-chengxin']['teaching']; cpf = self.ps['nju-chenpengfei']['teaching']
        self.assertEqual([t['level'] for t in cx], ['本科三年级','本科一年级','Level 3, Spring','Level 3, Fall'])
        self.assertEqual(cpf[1]['level'], '研究生一年级')
        for p in self.new:
            js = self.consumers()['people'][p['person_id']]
            tables = self.tables(p['person_id'])
            if p.get('teaching'):
                self.assertIn('公开教学事实',js['profileContent'])
                self.assertIn('不是本人求学学历',js['profileContent'])
                self.assertIn('公开教学事实',tables['总览'][1][6])
                for t in p['teaching']:
                    self.assertIn(t['course'], js['profileContent'])
                    self.assertIn(t['course'], tables['总览'][1][6])
                    self.assertIn(t['level'], js['profileContent'])
                    self.assertIn(t['level'], tables['总览'][1][6])
                    self.assertEqual(t['field_provenance']['note'], [t['source_id']])
                    if t.get('period'):
                        self.assertIn(t['period'], js['profileContent'])
                        self.assertIn(t['period'], tables['总览'][1][6])
                    self.assertNotIn(t['course'], js['graph'])
                    self.assertNotIn(t['course'], json.dumps(tables['经历'],ensure_ascii=False))
            else:
                self.assertNotIn('公开教学事实',js['profileContent'])
            self.assertNotIn('公开教学事实',js['lineage'])

    def test_17_structured_only_arrays_do_not_change_any_ui_graph_or_export_table(self):
        stripped = copy.deepcopy(self.d)
        for p in stripped['people']:
            if p['person_id'] in SOURCE_FINGERPRINTS:
                for key in ('academic_service','honors','research_projects'):
                    p.pop(key,None)
                self.assertEqual(self.tables(p['person_id'],p),self.tables(p['person_id']))
        # Metamorphic actual-app test: removing all three independent arrays
        # cannot change profile, lineage, graph, school pages, filters or roster.
        self.assertEqual(self.consumers(stripped),self.consumers())

    def test_18_exact_selected_order_scope_partial_status_and_real_organization(self):
        for school, ids in EXPECTED_IDS.items():
            batch = next(b for b in self.d['profile_batches'] if b['batch_id']==f'faculty-batch-{school}-rotation-002')
            self.assertEqual(batch['selection'],EXPECTED_SCOPES[school])
            self.assertEqual(batch['limitations'][:len(EXPECTED_LIMITATIONS[school])],EXPECTED_LIMITATIONS[school])
            self.assertEqual(batch['source_files'],[{'sha256':SOURCE_PACKAGE_SHA256[school]}])
            self.assertEqual(batch['people_count'],3)
            self.assertFalse(batch['coverage_complete']); self.assertFalse(batch['fields_complete'])
            entries = [e for e in self.r['entries'][1700:1706] if e['school_id']==school]
            self.assertEqual([e['source_person_id'] for e in entries],ids)
            self.assertEqual([e['name'] for e in entries],EXPECTED_NAMES[school])
            for entry in entries:
                self.assertEqual(entry['selection_basis'],EXPECTED_SCOPES[school])
                self.assertEqual(entry['source_evidence_status'],'official_roster_body_read')
                self.assertEqual(entry['campus_ids'],[])
        orgs = {o['id']:o for o in self.d['org_units']}
        self.assertEqual(orgs['hnu-chemistry-department']['parent_id'],'hnu-chemistry')
        self.assertEqual(orgs['hnu-chemistry-department']['type'],'department')
        self.assertEqual(orgs['nju-astronomy-space-science']['type'],'school')
        for p in self.new:
            self.assertEqual(p['campus_ids'],[]); self.assertIsNone(p['campus'])
            self.assertNotIn('professor',p['department_id'])

    def test_19_canonical_tables_keep_five_sheets_main_rows_sources_and_tristates(self):
        for p in self.new:
            tables = self.tables(p['person_id'])
            self.assertEqual(list(tables),['总览','经历','导师','论文','来源'])
            self.assertEqual(len(tables['经历']),len(p['education'])+len(p['appointments'])+1)
            self.assertEqual(len(tables['导师']),len(p['mentors'])+1)
            self.assertEqual(len(tables['论文']),len(p['publications'])+1)
            for e,row in zip(p['education'],tables['经历'][1:]):
                self.assertEqual(row[1:6],[e['event_kind'],e['degree'],e['institution'],e['start'],e['end']])
                self.assertIs(row[11],e['degree_awarded'])
                self.assertEqual(row[12],e['degree_awarded_date'])
                self.assertEqual(row[16],e['source_url'])
            for rows in tables.values():
                self.assertEqual(rows[0][-2:],['所属学校ID','所属学校'])
                self.assertTrue(all(row[-2]==p['school_id'] for row in rows[1:]))

    def test_20_new_department_workbooks_reopen_equal_every_canonical_cell_and_link(self):
        from openpyxl import load_workbook
        departments = sorted({p['department_id'] for p in self.new})
        self.assertEqual(departments,['hnu-chemistry-department','nju-astronomy-space-science'])
        for dept in departments:
            people = [p for p in self.d['people'] if p['department_id']==dept]
            tables = build.department_tables(people,self.d['updated_at'],self.d['institutions'])
            wb = load_workbook(ROOT / 'site/exports' / (dept+'.xlsx'))
            try:
                self.assertEqual(wb.sheetnames,[name for name,_ in tables])
                for name,rows in tables:
                    ws=wb[name]
                    self.assertEqual((ws.max_row,ws.max_column),(len(rows),len(rows[0])))
                    for ri,row in enumerate(rows,1):
                        for ci,value in enumerate(row,1):
                            expected=build.clean(value)
                            self.assertEqual(ws.cell(ri,ci).value, None if expected=='' else expected,(dept,name,ri,ci))
                            link=ws.cell(ri,ci).hyperlink
                            self.assertEqual(link.target if link else None,expected if isinstance(expected,str) and build.valid_url(expected) else None,(dept,name,ri,ci,'hyperlink'))
            finally:
                wb.close()

    def test_21_selected_publications_are_exact_metadata_not_inferred_publisher_checks(self):
        pubs=[p for person in self.new for p in person['publications']]
        self.assertEqual(len(pubs),8)
        self.assertEqual(sum(p['individual_article_link_verified'] is True for p in pubs),3)
        for p in pubs:
            self.assertEqual(p['type'],'journal_article')
            self.assertTrue(p['type_label']); self.assertTrue(p['link_status_label'])
        for person in self.new:
            if person['school_id']=='hnu':
                p=person['publications'][0]
                self.assertIsNone(p['doi']); self.assertFalse(p['individual_article_link_verified'])
                self.assertEqual(p['url'],p['source_url'])
                self.assertEqual(p['link_status'],'official_publication_list_only')

    def test_22_no_application_changes_or_new_ui_columns(self):
        # Exact old release SHA/token assertions; current runtime is checked separately.
        for path,digest in UNCHANGED_APPLICATION_SHA256.items():
            self.assertEqual(hashlib.sha256(historical_bytes(path, root=published_root())).hexdigest(),RELEASE26_DISPLAY_SHA256.get(path, LAYOUT_FIX_APPLICATION_SHA256.get(path, ROTATION003_APPLICATION_SHA256.get(path, DISPLAY_FIX_APPLICATION_SHA256.get(path,digest)))),path)

    def test_23_consumer_statistics_are_partial_and_not_full_school_population(self):
        js=self.consumers()
        for school in EXPECTED_IDS:
            self.assertEqual(js['schoolStats'][school]['people'],sum(p['school_id']==school for p in self.d['people']))
            self.assertEqual(js['schoolStats'][school]['status'],'部分收录')
            self.assertIn('统计不代表全校教师或学院总数',js['schoolPages'][school])
            self.assertIn('非完整名录',js['schoolPages'][school])

    def test_24_validators_consumers_and_exports_are_nonmutating(self):
        before=copy.deepcopy((self.d,self.r))
        self.assertTrue(build.validate(self.d)); self.assertTrue(validate_rosters(self.r,self.d))
        for p in self.new:self.tables(p['person_id'])
        self.assertTrue(self.consumers()['unchanged'])
        self.assertEqual((self.d,self.r),before)

    def test_25_source_guard_rejects_candidate_award_date_and_independent_fact_corruption(self):
        pid='nju-chengxin';shape,digest=SOURCE_FINGERPRINTS[pid]
        for mutation in ('award','candidate','date','teaching','service','honor','project','provenance'):
            p=source_view(self.ps[pid])
            if mutation=='award':p['education'][0]['degree_awarded']=True
            elif mutation=='candidate':p['pending_education_candidates'][0]['degree_awarded']=None
            elif mutation=='date':p['education'][1]['degree_awarded_date']='2012-06'
            elif mutation=='teaching':p['teaching'][2]['level']='研究生'
            elif mutation=='service':p['academic_service'][0]['role']='导师'
            elif mutation=='honor':p['honors'][0]['year']=2012
            elif mutation=='project':p['research_projects'][0]['role']='参与'
            else:p['education'][1]['field_provenance']['degree_awarded']='nju-cx-mps-cv'
            self.assertNotEqual(object_sha(project_source_shape(p,shape)),digest,mutation)
        p=source_view(self.ps[pid]);del p['pending_education_candidates']
        with self.assertRaises(KeyError):project_source_shape(p,shape)
        p=source_view(self.ps[pid]);p['mentors'].append(copy.deepcopy(p['mentors'][0]))
        with self.assertRaises(AssertionError):project_source_shape(p,shape)
        pid='nju-chenpengfei';shape,digest=SOURCE_FINGERPRINTS[pid]
        p=source_view(self.ps[pid]);p['education'][3]['start']='2001-07'
        self.assertNotEqual(object_sha(project_source_shape(p,shape)),digest)

    def test_26_recursive_closure_guard_rejects_dangling_sources_and_candidate_edges(self):
        for mutation in ('source','url','provenance','edge','candidate_filter'):
            p=copy.deepcopy(self.ps['nju-chengxin'])
            if mutation=='source':p['academic_service'][0]['source_id']='missing'
            elif mutation=='url':p['pending_education_candidates'][0]['source_url']='https://wrong.example.test/'
            elif mutation=='provenance':p['research_projects'][0]['field_provenance']['name']='missing'
            elif mutation=='edge':p['mentors'][0]['education_event_id']='unidentified-candidate'
            else:p['pending_education_candidates'][0]['include_in_education_filters']=True
            with self.assertRaises(AssertionError,msg=mutation):assert_source_closure(p)

    def test_27_public_additions_exclude_contacts_caches_private_paths_and_raw_bodies(self):
        forbidden={'public_raw_cache','local_evidence_cache','cache_path','raw_html','raw_pdf','pdf_bytes','raw_response','email','phone','mobile','birthday','date_of_birth','student_list','students'}
        def check(value):
            if isinstance(value,dict):
                self.assertFalse(forbidden.intersection(value),forbidden.intersection(value))
                for child in value.values():check(child)
            elif isinstance(value,list):
                for child in value:check(child)
            elif isinstance(value,str):
                for prefix in ('/workspace/','/home/agent/','file://'):
                    self.assertNotIn(prefix,value)
        check(self.new);check(list(self.entries.values()))


    def test_28_six_reviewed_roster_identities_route_to_profiles_not_employment_claims(self):
        js = self.consumers()['rosters']
        snapshot = release25_view(json.loads((ROOT / 'site/rosters.json').read_text()), 'site/rosters.json')
        built = {e['source_person_id']:e for e in snapshot['entries'] if e.get('source_person_id') in SOURCE_FINGERPRINTS}
        for p in self.new:
            pid=p['person_id']; entry=self.entries[pid]; evidence=entry['profile_identity_evidence']
            self.assertTrue(entry['profile_identity_verified'])
            self.assertEqual(entry['verified_person_id'],pid)
            self.assertEqual(entry['name'],p['name'])
            self.assertEqual(entry['school_id'],p['school_id'])
            self.assertEqual(entry['department_id'],p['department_id'])
            self.assertEqual(entry['official_profile_url'],p['official_url'])
            self.assertEqual(evidence['reviewed_source_sha256'],SOURCE_PACKAGE_SHA256[p['school_id']])
            self.assertEqual(evidence['independent_review_sha256'],REVIEW_SHA256[p['school_id']])
            self.assertEqual(evidence['source_ids'],p['identity_evidence']['source_ids'])
            self.assertIn('不表示学历完整、全职或人事在岗已核',evidence['scope_note'])
            self.assertTrue(js[pid]['matched'])
            self.assertEqual(js[pid]['status'],'身份已关联人物档案（任职状态另核）')
            self.assertIn('#/person/'+pid,js[pid]['html'])
            self.assertTrue(built[pid]['profile_identity_verified'])
            self.assertEqual(built[pid]['verified_person_id'],pid)
            self.assertEqual(built[pid]['profile_identity_evidence'],evidence)



    def test_29_research_fields_adapt_verbatim_in_source_order_to_profile_and_directory(self):
        from html import escape
        js=self.consumers()
        for p in self.new:
            names=[field['name'] for field in p['research_fields']]
            expected='；'.join(names)
            source_ids=list(dict.fromkeys(field['source_id'] for field in p['research_fields']))
            self.assertTrue(names)
            self.assertEqual(p['research_interests'],expected)
            self.assertEqual(p['research_interests_source_ids'],source_ids)
            self.assertEqual(p['professional_field_provenance']['research_interests'],source_ids)
            self.assertEqual(p['research_interests_derivation'],'verbatim_original_research_fields_names_in_source_order_for_existing_consumer')
            self.assertIn('<h3>研究方向</h3><p>'+escape(expected)+'</p>',js['people'][p['person_id']]['profileContent'])
            directory=js['directories'][p['school_id']]
            row=next(part for part in directory.split('<tr>') if '#/person/'+p['person_id']+'"' in part)
            self.assertIn('<small>'+escape(expected)+'</small>',row)
            self.assertNotIn('研究方向待补录',row)
            stripped=copy.deepcopy(p)
            for key in ('research_interests','research_interests_source_ids','research_interests_derivation'):
                stripped.pop(key)
            stripped['professional_field_provenance'].pop('research_interests')
            # Existing workbooks do not have a research column: this compatibility
            # alias changes website display only and cannot add/reclassify rows.
            self.assertEqual(self.tables(p['person_id'],stripped),self.tables(p['person_id']))



# Release25 exact compatibility overlay; original historical pins above remain unchanged.
ROTATION003_APPLICATION_SHA256 = {'site/app.js': '5d4efe8a3c85cbca0125b950a70d7494e3d34a2e67240e6043347903965c8d4b', 'site/model.js': '9dbc155288c1ee0a93fd5d11cd8a7c63724dc0bde6b117a2f71d327e34be3fef', 'site/rosters.js': 'e4bb3c53da7a4ff3c00c1109651d66cc479c20dc265e7705e41213e3019a3502', 'site/index.html': '3b5f7afdd4ac96aaca402884976bccb33cb9a1d2a1ad27793cf7cdd82e5970bf', 'scripts/build.py': '92b6ec03d18666503a7e890e09faffe2e3ac7be61b017ec85b291c1498addacb', 'scripts/build_rosters.py': '29c5dee611bfaef3fd6afd1df9480b254652c51b1c4ce1c9568d75cb799319fe'}


# Exact layout-only overlay; all original source/data/application pins remain above.
LAYOUT_FIX_APPLICATION_SHA256 = {'site/styles.css': '3b01ec6a5674982d048237e5426aa107bd622662d2758913c81bc6f5c4c0ceca', 'site/app.js': '5ce1dbc7161a048ab263cdff313bc904544b362c13d39cc5963886b409f037e2', 'site/rosters.js': '7536ef08c4032e05e1ad3c09532858dd42fd68c20968c82e2146481cf6d9c85f', 'site/index.html': '7e6331d4846c13d64f0a5b5b65a9d52539df537b2ce20ef26563669739f8f2a3'}

# Exact three-file heading/cache correction; all earlier golden constants remain.
RELEASE26_DISPLAY_SHA256={'site/index.html': 'a6e57e88a623157fd3103a6099b9e72b21a0435c9fa62530a41469814844242c', 'site/app.js': '19945c1392982e4b3a1df95d93c271e957bf16b2a238f8d63d24ed892c722dea', 'site/rosters.js': 'd060662fe29e2d5209ca7cabd89121f253c6df69878d776e652f2fe8f8eb45a9'}

if __name__ == '__main__':
    unittest.main()
