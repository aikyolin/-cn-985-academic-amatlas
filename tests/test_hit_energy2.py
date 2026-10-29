"""Exact HIT two-person additive checks; only the two authorized generated XLSX are read."""
import copy,hashlib,json,os,sys,unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from atlas_storage import load_atlas
import build,build_rosters
from openpyxl import load_workbook
SHA='5a41bb46ed311be7754576b313669f67626f5b44b0f1bf9f9c92fde4fe67ad9e'
IDS=['hit-zhaozhiyuan','hit-yanhan-energy']
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
class HitEnergyTwoTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.parent=Path(os.environ['HIT_ENERGY_BASELINE_ROOT']).resolve()
  cls.receipt_path=Path(os.environ['HIT_ENERGY_BUILD_RECEIPT']).resolve()
  if cls.parent==ROOT or cls.receipt_path.is_relative_to(ROOT):raise ValueError('separate immutable baseline/external receipt required')
  fp=ROOT/'tests/fixtures/hit-energy2-source-mapping.json'
  if sha(fp)!=SHA:raise ValueError('fixed mapping SHA mismatch')
  cls.fixed=json.loads(fp.read_text());cls.base=load_atlas(cls.parent/'data/atlas.json');cls.actual=load_atlas(ROOT/'data/atlas.json');cls.site=load_atlas(ROOT/'site/atlas.json');cls.receipt=json.loads(cls.receipt_path.read_text())
 def test_01_old_1149_whole_objects_and_all_other_root_values_are_preserved(self):
  self.assertEqual(hashlib.sha256(json.dumps(self.base,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest(),'152fcd92db9dc7f2d5fc008591d42b6d85c77d195753b21627e632bb7e06671c');self.assertEqual(len(self.base['people']),1149);self.assertEqual(len(self.actual['people']),1151)
  self.assertEqual(self.actual['people'][:1149],self.base['people']);self.assertEqual(self.actual['people'][1149:],self.fixed['people'])
  self.assertEqual([p['person_id'] for p in self.fixed['people']],IDS)
  self.assertEqual(self.actual['profile_batches'],self.base['profile_batches']+self.fixed['profile_batches'])
  restored=copy.deepcopy(self.actual);restored['people']=restored['people'][:1149];restored['profile_batches']=restored['profile_batches'][:-1]
  self.assertEqual(restored,self.base);self.assertEqual(len(self.actual['pending_people']),84)
  self.assertEqual((ROOT/'data/rosters.json').read_bytes(),(self.parent/'data/rosters.json').read_bytes())
  for k,v in self.actual.items():self.assertEqual(self.site[k],v,k)
  others=[p for p in (ROOT/'site/people').glob('*.json') if p.name!='hit.json'];self.assertEqual(len(others),38)
  for p in others:self.assertEqual(p.read_bytes(),(self.parent/'site/people'/p.name).read_bytes(),p.name)
 def test_02_source_semantics_and_negative_projection_changes(self):
  people=self.actual['people'][1149:];events=[e for p in people for e in p['education']]
  self.assertEqual(len(events),7);self.assertEqual([sum(e['degree_awarded'] is v for e in events)for v in (True,None,False)],[2,3,2])
  self.assertEqual(sum(len(p['appointments'])for p in people),4);self.assertEqual(sum(len(p['mentors'])for p in people),3);self.assertEqual(sum(len(p['publications'])for p in people),2)
  for p in people:
   self.assertEqual(p['school_id'],'hit');self.assertEqual(p['department_id'],'hit-energy');self.assertEqual(p['campus_ids'],['harbin'])
   for e in p['education']:self.assertIn('degree_awarded_date',e);self.assertIsNone(e['degree_awarded_date'])
   for m in p['mentors']:self.assertEqual(m['employment_status'],'unknown');self.assertIsNone(m['person_id']);self.assertIsNone(m['formal_dissertation_advisor_verified'])
  changed=copy.deepcopy(people);changed[0]['education'][0]['degree_awarded']=False
  self.assertNotEqual(changed,self.fixed['people'])
  changed=copy.deepcopy(people);changed[1]['mentors'][2]['education_event_id']='hit-yanhan-energy-education-2'
  self.assertNotEqual(changed,self.fixed['people'])
 def assert_tables(self,path,tables):
  book=load_workbook(path,read_only=True,data_only=False)
  try:
   self.assertEqual(book.sheetnames,list(tables))
   for name,rows in tables.items():
    expected=[[build.clean(v) if build.clean(v)!='' else None for v in row]for row in rows]
    width=max(map(len,expected),default=0)
    expected=[row+[None]*(width-len(row)) for row in expected]
    actual=[list(row)for row in book[name].iter_rows(values_only=True)]
    self.assertEqual(actual,expected,name)
  finally:book.close()
 def test_03_only_hit_energy_five_tables_and_global_roster_nine_tables(self):
  members=[p for p in self.actual['people']if p['department_id']=='hit-energy'];self.assertEqual([p['person_id']for p in members],IDS)
  self.assertFalse(any(p['department_id']=='hit-energy'for p in self.base['people']))
  tables=dict(build.department_tables(members,self.actual['updated_at'],self.actual['institutions']));self.assertEqual(list(tables),['总览','经历','导师','论文','来源'])
  self.assert_tables(ROOT/'site/exports/hit-energy.xlsx',tables)
  rosters=json.loads((ROOT/'data/rosters.json').read_text());rt=build_rosters.roster_tables(build_rosters.prepare_rosters(rosters,self.actual),self.actual);self.assertEqual(len(rt),9)
  self.assert_tables(ROOT/'site/exports/待核对名录.xlsx',rt)
  # These exact source-derived checks do not rely solely on shared table functions.
  self.assertEqual([r[0]for r in tables['总览'][1:]],[p['name']for p in members]);self.assertEqual(len(tables['导师'])-1,3);self.assertEqual(len(tables['论文'])-1,2)
  headers=tables['经历'][0];self.assertEqual(headers[11],'来源记载学位');self.assertEqual(headers[14],'备注及冲突')
  self.assertEqual(len(tables['经历'])-1,11)
  for p in members:
   rows=[r for r in tables['经历'][1:]if r[0]==p['name']];self.assertEqual(len(rows),len(p['education'])+len(p['appointments']))
   for row,e in zip(rows[:len(p['education'])],p['education'],strict=True):
    self.assertEqual(row[headers.index('事件类型')],e['event_kind']);self.assertEqual(row[headers.index('阶段/职位')],e['degree']);self.assertIs(row[headers.index('来源记载学位')],e['degree_awarded']);self.assertEqual(row[headers.index('开始')],e['start']);self.assertEqual(row[headers.index('结束')],e['end']);self.assertEqual(row[headers.index('备注及冲突')],e['note'])
   for row,a in zip(rows[len(p['education']):],p['appointments'],strict=True):
    self.assertEqual(row[headers.index('事件类型')],'employment');self.assertEqual(row[headers.index('阶段/职位')],a['title']);self.assertEqual(row[headers.index('开始')],a['start']);self.assertEqual(row[headers.index('结束')],a['end']);self.assertEqual(row[headers.index('备注及冲突')],a['note']);self.assertIsNone(row[headers.index('来源记载学位')])
 def test_04_receipt_exact_two_workbooks_and_complete_current_output_hashes(self):
  r=self.receipt;self.assertEqual(r['status'],'HIT_ENERGY2_RESTRICTED_BUILD_PASS');self.assertEqual(r['base_commit'],'bd86a1da000289a132aad5d894b5242142f55241');self.assertEqual(r['build_id'],self.site['build_id'])
  workbooks={'site/exports/hit-energy.xlsx','site/exports/待核对名录.xlsx'}
  outputs=workbooks|{'site/people/'+p.name for p in (ROOT/'data/people').glob('*.json')}|{'site/atlas.json','site/rosters.json','site/exports/待核对名录.csv'}
  self.assertEqual(len(outputs),44);self.assertEqual(set(r['generated_workbooks']),workbooks);self.assertEqual(set(r['generated_output_inventory']),outputs);self.assertEqual(set(r['generated_sha256']),outputs)
  for rel in outputs:self.assertEqual(sha(ROOT/rel),r['generated_sha256'][rel],rel)
  self.assertEqual(r['old_full_suite'],'NOT_RUN');self.assertFalse(any('xmu-'in p for p in workbooks))
if __name__=='__main__':unittest.main()
