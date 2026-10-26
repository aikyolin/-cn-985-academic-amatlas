from atlas_test_io import read_live_json
"""Current candidate checks, deliberately NOT run against historical fixtures.

The fixed published release is only the independent conservation reference.
"""
import copy
import hashlib
import json
from pathlib import Path
import sys
import unittest
from published_fixture import published_root
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import build
from build_rosters import validate_rosters
from public_metadata import check_public_tree

MAPPING_SHA256 = 'b08daade03c91c774ecda2bc9469f5d92497be873732c5bfb58c59ab9a244475'
FIRST_IDS = ['nju-shiyini', 'nju-yangderui']
NEXT_MAPPING_SHA256 = '6b7347159e025485837de10750b40e7af329f229db559e273f4f9145095a6079'
SEU_MAPPING_SHA256 = 'e628de5b926bc65288b908aa2c8fc85fd5ff55f156891e2ab573acacaaa86795'
SCU_MAPPING_SHA256 = 'c865a6548faf84d525bf2d42c2f6df1584a1e27e6a4298f845180fe3c656addf'
UESTC_MAPPING_SHA256 = '3237a4596f84af20c149ac69f9e30917fe7f7df31fce021c5005a11cab18d1c3'
HNU_MAPPING_SHA256 = '5415b445dde9ab09dd38a97fc0f1d0136c08b2fa8c2667d32ee35b4301dff4c6'
BNU_MAPPING_SHA256 = '9c3004d31ae1327e5c7ffec1251fc71be3b6417606b0fdc6fe583f6d7e4b36d3'
TJU_MAPPING_SHA256 = 'f82df77db84d07c64b594b46047ed775c45a1001734a06493035094e2473bc1f'
NKU_MAPPING_SHA256 = '6b999966d28b6396f31074ececc75c7e152e1dcedcd22a3d7b77b82ccde41aae'
DLUT_MAPPING_SHA256 = 'a579968911068d28b5371ff96d3ba4557d18725d0401d32ae236df092f359909'
JLU_MAPPING_SHA256 = 'f010fb82e3ad522830c7c79b49fc06a0042e1e47daf1d424d3c9359641dc4856'
SDU_MAPPING_SHA256 = 'ea4ea15f0fa0a691d113aac302779bde3aa2be5589d6a5b33751006581619743'
CSU_MAPPING_SHA256 = '0c227fdedfe70344f92db0fb32e905a57e60979fb28f35a00a70a29effe76609'
HUST_MAPPING_SHA256 = '5cccae4a6117af0caec63da48e6fd196f393c30090e2f203d3ab7ae0e0f8dff0'
WHU_MAPPING_SHA256 = '5b2a4e707631c45919bd8aa8b711285e2699989b630dd50148c8002c2c558e4c'
XJTU_MAPPING_SHA256 = 'e5a81ca95e85f6cc81e7fb4a16984dfecdf1e8089538cd6b315dcafeb09c369d'
RUC_MAPPING_SHA256 = '39d6b531ba6c3e4a72ec1f6a6f8173b6582dac29f047187949556e55fae1683c'
ECNU_MAPPING_SHA256 = 'b35c42f95987f63bdd18773349270e8f7b289965f7923519776a9cc89b64c609'
BUAA_MAPPING_SHA256 = '8bbe4870f6c43131d5f4c9ca08dad0f25940cbb001491b3f3151435c96d07f24'
TONGJI_MAPPING_SHA256 = 'f337b168e9987b41d07704cbbbdc34ff904cece594eabe1551b625bf0614e09d'
BIT_MAPPING_SHA256 = 'bdb9ad4df6079b40b8b2d5ff0038aa3536abda806ef1feb42d73a1178d35a7a4'
NWPU_MAPPING_SHA256 = '5a569c0898261557784e824a3f49c3fd9f29850ae2d6959ee2d33070d38ce5b8'
SCUT_MAPPING_SHA256 = '7198a55982b7e8926997a1e52963c35a420aa9458e6e2ca2312a818c1de49d80'
SCUT_IDS = ['scut-chenguangxu', 'scut-heping', 'scut-chewenquan', 'scut-caoyufei', 'scut-zhongzhiheng', 'scut-yangxiongwen', 'scut-yangwei', 'scut-liuxiaoni', 'scut-guodifan', 'scut-qilili', 'scut-zhongyouliang-mathematics']
NWPU_IDS = ['nwpu-chenkaijie', 'nwpu-heyibai', 'nwpu-dai-yuchao', 'nwpu-zhao-ruiqin', 'nwpu-songdingbonan', 'nwpu-zhaoxueai', 'nwpu-zhangyi', 'nwpu-fupengyu', 'nwpu-zhengye', 'nwpu-sunhao-applied-statistics', 'nwpu-yangkunde', 'nwpu-sunxiaodi']
BIT_IDS = ['bit-songtian', 'bit-shifugui', 'bit-guanlei', 'bit-xiaenjun', 'bit-shanjiefei', 'bit-tongxuguang', 'bit-sunmiao', 'bit-lumingfei', 'bit-yangkuan', 'bit-liuning', 'bit-wanglang', 'bit-xuejin']
TONGJI_IDS = ['tongji-liuhuajie', 'tongji-yangjie', 'tongji-zhaoqianfan', 'tongji-wenxianghao', 'tongji-huding', 'tongji-yangweitao', 'tongji-gexin', 'tongji-leikewen', 'tongji-piyong', 'tongji-xielihua', 'tongji-huominfeng', 'tongji-sunxiaoting']
BUAA_IDS = ['buaa-quanquan', 'buaa-daiwei', 'buaa-dingchuan', 'buaa-zhou-xuefeng', 'buaa-tang-bin', 'buaa-jinhao', 'buaa-niuchaohui', 'buaa-wangtianhua', 'buaa-wuruilin', 'buaa-chenyan-sportstraining', 'buaa-zhouronggang-engineeringpsychology', 'buaa-yuehanwei']
ECNU_IDS = ['ecnu-yinyaling', 'ecnu-guanshuguang', 'ecnu-jiangxuefeng', 'ecnu-wangxuena', 'ecnu-liangchaoqun', 'ecnu-huangchenxi', 'ecnu-yinzhihua', 'ecnu-xijuzhe', 'ecnu-xusheng', 'ecnu-yinguoming', 'ecnu-zhuhu-art', 'ecnu-gejun']
RUC_IDS = ['ruc-yaohaifang', 'ruc-liu-danfeng', 'ruc-qitianjiao', 'ruc-xuwangli', 'ruc-lijingping', 'ruc-wangshuixiong', 'ruc-zhaoyandong', 'ruc-anyaoshun']
XJTU_IDS = ['xjtu-fenggenfu', 'xjtu-cui-chenhui', 'xjtu-shan-wenhua', 'xjtu-xu-yubin', 'xjtu-luojiu', 'xjtu-yao-xiuying-wave168', 'xjtu-xuhaoran-chinese', 'xjtu-shangyufan-organization', 'xjtu-yuyan', 'xjtu-peileilei', 'xjtu-zhoumi-organization', 'xjtu-zhuzhengwei']
WHU_IDS = ['whu-yufeng', 'whu-zhaolin', 'whu-wangliuquan', 'whu-wanghaowu', 'whu-weibin', 'whu-mazhaocheng', 'whu-miaoxiaoping', 'whu-yuanbifeng-publichealth', 'whu-yanhong-publichealth']
HUST_IDS = ['hust-baoyukun', 'hust-baiyun', 'hust-changyaping', 'hust-hongmin', 'hust-chenxianhong', 'hust-liuchengbin', 'hust-kelan', 'hust-xieruixiang', 'hust-liwei-education', 'hust-longlirong', 'hust-dengshiming', 'hust-zhangyanglei']
CSU_IDS = ['csu-zhangligang', 'csu-daiyu', 'csu-yanlei', 'csu-wangmiao', 'csu-fenghui', 'csu-denglu', 'csu-yiqiao', 'csu-guocan', 'csu-wenyu', 'csu-yaohailin', 'csu-yanjiexiong', 'csu-fengjie']
SDU_IDS = ['sdu-zhangna', 'sdu-wangliping', 'sdu-zhanghaiyan', 'sdu-sunlixin', 'sdu-yidonglan', 'sdu-wangkefang', 'sdu-liyuli', 'sdu-fanghui', 'sdu-lijing-nursing', 'sdu-lichaofan-policy', 'sdu-luhaobin', 'sdu-liusenlin']
JLU_IDS = ['jlu-anni', 'jlu-anshengji', 'jlu-wangting-economics', 'jlu-lvxuanru', 'jlu-liwenjun', 'jlu-guochunfang', 'jlu-zhouxin', 'jlu-wangjuan', 'jlu-jinlina', 'jlu-yuxianfeng', 'jlu-sunwanchun', 'jlu-wangdehui']
DLUT_IDS = ['dlut-zhangjialiang', 'dlut-zhangrui', 'dlut-cuimiao', 'dlut-wanghuanming', 'dlut-zhangwei1', 'dlut-liulinlin-law', 'dlut-suyongjian-education', 'dlut-liumengyang', 'dlut-zhouwenjie-art', 'dlut-lichong-higher-education', 'dlut-lijunfeng', 'dlut-mahui']
NKU_IDS = ['nku-lingfeng', 'nku-zhangjiayu-philosophy', 'nku-luzhaoyong', 'nku-xuexue', 'nku-yangzhenhong', 'nku-zhangzhihong', 'nku-caishengxia', 'nku-maodaqing-medical', 'nku-guanyongqiang-economic-history', 'nankai-haolan', 'nankai-luozhenya', 'nankai-liulili']
TJU_IDS = ['tju-xuchunjian', 'tju-zhaobohang', 'tju-wenxiaotong', 'tju-mawenchao', 'tju-shizhiyuan', 'tju-tanhao', 'tju-zhaokuo', 'tju-chentong', 'tju-liuconghong', 'tju-zhaoyanjun-pharmaceutics', 'tju-daisong', 'tju-hanzhongjie']
BNU_IDS = ['bnu-hanxue', 'bnu-hufengqin', 'bnu-liushaomin', 'bnu-sudianna', 'bnu-wangchen-music', 'bnu-guo-jia-wave169', 'bnu-wang-cixiao-wave190', 'bnu-zeng-xianglong-wave190', 'bnu-lichong', 'bnu-luoxinhui', 'bnu-nitengda', 'bnu-zhanglu']
HNU_IDS = ['hnu-huyueqiang', 'hnu-liujie', 'hnu-libing-education-history', 'hnu-hubicheng', 'hnu-yuanxiang-design', 'hnu-yuanzhaohui-architecture', 'hnu-yanjianfei', 'hnu-sunqingjuan', 'hnu-tanhaibo', 'hnu-huying']
UESTC_IDS = ['uestc-shangjizheng', 'uestc-shenhua', 'uestc-fangjiaming', 'uestc-kangjian', 'uestc-qianhao', 'uestc-liangjian', 'uestc-nidebing', 'uestc-dengchenliang', 'uestc-zhouguanyu', 'uestc-zhupeiyong']
SCU_IDS = ['scu-baishibing', 'scu-baojianjun', 'scu-caoya', 'scu-longbingwei', 'scu-fanzhenxin', 'scu-xuliang', 'scu-fengxianqiong', 'scu-penghua-classics', 'scu-lifei', 'scu-huangzongxian', 'scu-wuyongqiang']
SEU_IDS = ['seu-dinghui', 'seu-sushu', 'seu-zhangkaifeng', 'seu-zhangjianqiong', 'seu-liupeidang', 'seu-shan-pingji-wave171', 'seu-he-shaojin-wave171', 'seu-zhangxueyi', 'seu-dongqun', 'seu-longdiyong']
NEXT_IDS = ['nju-huxiaowu', 'nju-caipingqiang', 'nju-linzhaoyu', 'nju-huying', 'nju-zhangwei-linguistics', 'nju-lvlinhai-education', 'nju-liyougen-law', 'nju-chenbaojun', 'nju-jiangfei', 'nju-zhoujiaxin']

def read(root, rel):
    return (json.loads((root / rel).read_text()) if root == published_root() else read_live_json(root / rel))

def mapping():
    raw = (ROOT / 'tests/fixtures/independent-nju-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == MAPPING_SHA256, 'Reviewed new-source mapping changed'
    return json.loads(raw)

def next_mapping():
    raw = (ROOT / 'tests/fixtures/independent-nju10-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == NEXT_MAPPING_SHA256, 'Reviewed new ten-person mapping changed'
    return json.loads(raw)

def seu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-seu10-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == SEU_MAPPING_SHA256, 'Reviewed SEU ten-person mapping changed'
    return json.loads(raw)

def scu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-scu11-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == SCU_MAPPING_SHA256, 'Reviewed SCU eleven-person mapping changed'
    return json.loads(raw)

def uestc_mapping():
    raw = (ROOT / 'tests/fixtures/independent-uestc10-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == UESTC_MAPPING_SHA256, 'Reviewed UESTC ten-person mapping changed'
    return json.loads(raw)

def hnu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-hnu10-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == HNU_MAPPING_SHA256, 'Reviewed HNU ten-person mapping changed'
    return json.loads(raw)

def bnu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-bnu12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == BNU_MAPPING_SHA256, 'Reviewed BNU twelve-person mapping changed'
    return json.loads(raw)

def tju_mapping():
    raw = (ROOT / 'tests/fixtures/independent-tju12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == TJU_MAPPING_SHA256, 'Reviewed TJU twelve-person mapping changed'
    return json.loads(raw)

def nku_mapping():
    raw = (ROOT / 'tests/fixtures/independent-nku12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == NKU_MAPPING_SHA256, 'Reviewed NKU twelve-person mapping changed'
    return json.loads(raw)

def dlut_mapping():
    raw = (ROOT / 'tests/fixtures/independent-dlut12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == DLUT_MAPPING_SHA256, 'Reviewed DLUT twelve-person mapping changed'
    return json.loads(raw)

def jlu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-jlu12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == JLU_MAPPING_SHA256, 'Reviewed JLU twelve-person mapping changed'
    return json.loads(raw)

def sdu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-sdu12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == SDU_MAPPING_SHA256, 'Reviewed SDU twelve-person mapping changed'
    return json.loads(raw)

def csu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-csu12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == CSU_MAPPING_SHA256, 'Reviewed CSU twelve-person mapping changed'
    return json.loads(raw)

def hust_mapping():
    raw = (ROOT / 'tests/fixtures/independent-hust12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == HUST_MAPPING_SHA256, 'Reviewed HUST twelve-person mapping changed'
    return json.loads(raw)

def whu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-whu9-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == WHU_MAPPING_SHA256, 'Reviewed WHU nine-person mapping changed'
    return json.loads(raw)

def xjtu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-xjtu12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == XJTU_MAPPING_SHA256, 'Reviewed XJTU twelve-person mapping changed'
    return json.loads(raw)

def ruc_mapping():
    raw = (ROOT / 'tests/fixtures/independent-ruc8-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == RUC_MAPPING_SHA256, 'Reviewed RUC eight-person mapping changed'
    return json.loads(raw)

def ecnu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-ecnu12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == ECNU_MAPPING_SHA256, 'Reviewed independent ECNU twelve-person mapping changed'
    return json.loads(raw)

def buaa_mapping():
    raw = (ROOT / 'tests/fixtures/independent-buaa12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == BUAA_MAPPING_SHA256, 'Reviewed independent BUAA twelve-person mapping changed'
    return json.loads(raw)

def tongji_mapping():
    raw = (ROOT / 'tests/fixtures/independent-tongji12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == TONGJI_MAPPING_SHA256, 'Reviewed independent Tongji twelve-person mapping changed'
    return json.loads(raw)

def bit_mapping():
    raw = (ROOT / 'tests/fixtures/independent-bit12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == BIT_MAPPING_SHA256, 'Reviewed independent BIT twelve-person mapping changed'
    return json.loads(raw)

def nwpu_mapping():
    raw = (ROOT / 'tests/fixtures/independent-nwpu12-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == NWPU_MAPPING_SHA256, 'Reviewed independent NWPU twelve-person mapping changed'
    return json.loads(raw)

def scut_mapping():
    raw = (ROOT / 'tests/fixtures/independent-scut11-source-mapping.json').read_bytes()
    assert hashlib.sha256(raw).hexdigest() == SCUT_MAPPING_SHA256, 'Reviewed independent SCUT eleven-person mapping changed'
    return json.loads(raw)

def reviewed_additions():
    first, second, third, fourth, fifth, sixth, seventh, eighth, ninth, tenth, eleventh, twelfth, thirteenth, fourteenth, fifteenth, sixteenth, seventeenth, eighteenth, nineteenth, twentieth, twenty_first, twenty_second, twenty_third = mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping(), ruc_mapping(), ecnu_mapping(), buaa_mapping(), tongji_mapping(), bit_mapping(), nwpu_mapping(), scut_mapping()
    assert [p['person_id'] for p in first['people']] == FIRST_IDS
    assert [p['person_id'] for p in second['people']] == NEXT_IDS
    assert [p['person_id'] for p in third['people']] == SEU_IDS
    assert [p['person_id'] for p in fourth['people']] == SCU_IDS
    assert [p['person_id'] for p in fifth['people']] == UESTC_IDS
    assert [p['person_id'] for p in sixth['people']] == HNU_IDS
    assert [p['person_id'] for p in seventh['people']] == BNU_IDS
    assert [p['person_id'] for p in eighth['people']] == TJU_IDS
    assert [p['person_id'] for p in ninth['people']] == NKU_IDS
    assert [p['person_id'] for p in tenth['people']] == DLUT_IDS
    assert [p['person_id'] for p in eleventh['people']] == JLU_IDS
    assert [p['person_id'] for p in twelfth['people']] == SDU_IDS
    assert [p['person_id'] for p in thirteenth['people']] == CSU_IDS
    assert [p['person_id'] for p in fourteenth['people']] == HUST_IDS
    assert [p['person_id'] for p in fifteenth['people']] == WHU_IDS
    assert [p['person_id'] for p in sixteenth['people']] == XJTU_IDS
    assert [p['person_id'] for p in seventeenth['people']] == RUC_IDS
    assert [p['person_id'] for p in eighteenth['people']] == ECNU_IDS
    assert [p['person_id'] for p in nineteenth['people']] == BUAA_IDS
    assert [p['person_id'] for p in twentieth['people']] == TONGJI_IDS
    assert [p['person_id'] for p in twenty_first['people']] == BIT_IDS
    assert [p['person_id'] for p in twenty_second['people']] == NWPU_IDS
    assert [p['person_id'] for p in twenty_third['people']] == SCUT_IDS
    return {key: first[key] + second[key] + third[key] + fourth[key] + fifth[key] + sixth[key] + seventh[key] + eighth[key] + ninth[key] + tenth[key] + eleventh[key] + twelfth[key] + thirteenth[key] + fourteenth[key] + fifteenth[key] + sixteenth[key] + seventeenth[key] + eighteenth[key] + nineteenth[key] + twentieth[key] + twenty_first[key] + twenty_second[key] + twenty_third[key] for key in ('people', 'org_units', 'profile_batches')}

def validate_append(current, previous, additions):
    assert set(current) == set(previous), 'Unlisted canonical root fields'
    for key, value in previous.items():
        if key in ('people', 'org_units', 'profile_batches'):
            assert current[key] == value + additions[key], 'Historical or new source fact changed: ' + key
        else:
            assert current[key] == value, 'Unlisted change: ' + key
    assert [x['person_id'] for x in current['people'][len(previous['people']):]] == FIRST_IDS + NEXT_IDS + SEU_IDS + SCU_IDS + UESTC_IDS + HNU_IDS + BNU_IDS + TJU_IDS + NKU_IDS + DLUT_IDS + JLU_IDS + SDU_IDS + CSU_IDS + HUST_IDS + WHU_IDS + XJTU_IDS + RUC_IDS + ECNU_IDS + BUAA_IDS + TONGJI_IDS + BIT_IDS + NWPU_IDS + SCUT_IDS
    assert len({p['person_id'] for p in current['people']}) == len(current['people'])

class CurrentRelease(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.previous = read(published_root(), 'data/atlas.json')
        cls.actual = read(ROOT, 'data/atlas.json')
        cls.additions = reviewed_additions()

    def test_current_canonical_exact_append_conservation(self):
        validate_append(self.actual, self.previous, self.additions)
        self.assertEqual(len(self.actual['people']), len(self.previous['people']) + len(self.additions['people']))
        self.assertEqual(len(self.actual['pending_people']), 84)

    def test_published_614_prefix_is_independent_fixed_612_plus_first_mapping(self):
        first = mapping()
        self.assertEqual(len(self.previous['people']) + len(first['people']), 614)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + first[key]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_624_prefix_is_independent_fixed_612_plus_reviewed_nju_mappings(self):
        first, second = mapping(), next_mapping()
        self.assertEqual(len(self.previous['people']) + len(first['people']) + len(second['people']), 624)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + first[key] + second[key]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_634_prefix_is_independent_fixed_612_plus_reviewed_previous_mappings(self):
        first, second, third = mapping(), next_mapping(), seu_mapping()
        self.assertEqual(len(self.previous['people']) + len(first['people']) + len(second['people']) + len(third['people']), 634)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + first[key] + second[key] + third[key]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_645_prefix_preserves_independent_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 645)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_655_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 655)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_665_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 665)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_677_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 677)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_689_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 689)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_701_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 701)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_713_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 713)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_725_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 725)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_737_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 737)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_749_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 749)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_761_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 761)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_770_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 770)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_782_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 782)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_790_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping(), ruc_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 790)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_802_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping(), ruc_mapping(), ecnu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 802)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_814_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping(), ruc_mapping(), ecnu_mapping(), buaa_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 814)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_826_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping(), ruc_mapping(), ecnu_mapping(), buaa_mapping(), tongji_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 826)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_838_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping(), ruc_mapping(), ecnu_mapping(), buaa_mapping(), tongji_mapping(), bit_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 838)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_published_850_prefix_preserves_all_previous_source_mappings(self):
        previous_mappings = (mapping(), next_mapping(), seu_mapping(), scu_mapping(), uestc_mapping(), hnu_mapping(), bnu_mapping(), tju_mapping(), nku_mapping(), dlut_mapping(), jlu_mapping(), sdu_mapping(), csu_mapping(), hust_mapping(), whu_mapping(), xjtu_mapping(), ruc_mapping(), ecnu_mapping(), buaa_mapping(), tongji_mapping(), bit_mapping(), nwpu_mapping())
        self.assertEqual(len(self.previous['people']) + sum(len(m['people']) for m in previous_mappings), 850)
        for key in ('people', 'org_units', 'profile_batches'):
            expected = self.previous[key] + [item for m in previous_mappings for item in m[key]]
            self.assertEqual(self.actual[key][:len(expected)], expected, key)

    def test_old_fact_deletion_reordering_and_unknown_promotions_fail(self):
        # Mutate each independently protected era, not just whichever person is last.
        for index in (0, 612, 614, 624, 634, 645, 655, 665, 677, 689, 701, 713, 725, 737, 749, 761, 770, 782, 790, 802, 814, 826, 838, 850):
            for key, value in [('name', 'invented'), ('quick_link_eligible', True)]:
                candidate = copy.deepcopy(self.actual)
                if candidate['people'][index].get(key) == value:
                    value = 'invented replacement'
                candidate['people'][index][key] = value
                with self.subTest(index=index, key=key), self.assertRaises(AssertionError):
                    validate_append(candidate, self.previous, self.additions)
        mutations = [
            lambda d: d['people'].pop(0),
            lambda d: d['people'].reverse(),
            lambda d: d['people'][612]['education'][0].update(degree_awarded=False),
            lambda d: d['people'][614]['sources'][0].update(url='https://example.test/unapproved'),
            lambda d: d['people'].append(copy.deepcopy(d['people'][-1])),
            lambda d: d.update(unapproved=True),
            lambda d: d['pending_people'].clear(),
        ]
        for change in mutations:
            candidate = copy.deepcopy(self.actual); change(candidate)
            with self.subTest(change=change), self.assertRaises(AssertionError):
                validate_append(candidate, self.previous, self.additions)

    def test_current_schema_build_and_source_metadata(self):
        import jsonschema
        validator = jsonschema.Draft202012Validator(read(ROOT, 'schema/atlas.schema.json'), format_checker=jsonschema.FormatChecker())
        validator.validate(self.actual)
        self.assertTrue(build.validate(self.actual))
        self.assertTrue(validate_rosters(read(ROOT, 'data/rosters.json'), self.actual))
        check_public_tree(ROOT, phase='final', include_workbooks=True)

    def test_generated_live_data_is_actual_canonical_and_build_id(self):
        site = read(ROOT, 'site/atlas.json')
        for key, value in self.actual.items():
            self.assertEqual(site[key], value, key)
        from atlas_storage import storage_inputs, compute_build_id
        # Independent implementation of the reviewed versioned length-framed contract.
        inputs = [('data/' + label, raw) for label, raw in storage_inputs(ROOT / 'data/atlas.json')]
        inputs += [(rel, (ROOT / rel).read_bytes()) for rel in (
            'data/rosters.json', 'scripts/build.py', 'scripts/build_rosters.py',
            'scripts/deterministic_xlsx.py', 'scripts/atlas_storage.py', 'site/atlas-loader.js')]
        digest = hashlib.sha256(b'academic-atlas-build-school-parts-v1\0')
        for label, raw in inputs:
            name = label.encode('utf-8')
            digest.update(len(name).to_bytes(8, 'big')); digest.update(name)
            digest.update(len(raw).to_bytes(8, 'big')); digest.update(raw)
        expected = digest.hexdigest()[:16]
        self.assertEqual(compute_build_id(ROOT), expected)
        self.assertEqual(site['build_id'], expected)
        self.assertEqual(read(ROOT, 'site/rosters.json')['build_id'], expected)
        self.assertEqual(site['exports'], build.department_exports(self.actual))
        new = next(x for x in site['exports'] if x['department_id'] == 'nju-social-sciences')
        expected_members = [p for p in self.additions['people'] if p['department_id'] == 'nju-social-sciences']
        self.assertEqual(new['people_count'], len(expected_members))

    def test_unaffected_old_workbooks_and_roster_facts_are_unchanged(self):
        old = published_root()
        affected = {p['department_id'] for p in self.additions['people']}
        for file in (old / 'site/exports').glob('*'):
            if file.name == '待核对名录.xlsx':
                continue  # Sole global count cell is tested separately below.
            if file.stem in affected:
                continue  # Changed membership: full old/new source-derived tables tested below.
            self.assertEqual((ROOT / 'site/exports' / file.name).read_bytes(), file.read_bytes(), file.name)
        self.assertEqual((ROOT / 'data/rosters.json').read_bytes(), (old / 'data/rosters.json').read_bytes())
        roster = read(ROOT, 'site/rosters.json'); previous = read(old, 'site/rosters.json')
        roster.pop('build_id'); previous.pop('build_id'); self.assertEqual(roster, previous)

    def test_roster_workbook_changes_only_global_profile_count_cell(self):
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        rel = 'site/exports/待核对名录.xlsx'
        member = 'xl/worksheets/sheet2.xml'
        with ZipFile(published_root() / rel) as old, ZipFile(ROOT / rel) as new:
            self.assertEqual(old.namelist(), new.namelist())
            for name in old.namelist():
                before, after = old.read(name), new.read(name)
                if name != member:
                    self.assertEqual(before, after, name)
                else:
                    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    a = ET.fromstring(before).find('.//s:c[@r="A39"]', ns)
                    b = ET.fromstring(after).find('.//s:c[@r="A39"]', ns)
                    self.assertIsNotNone(a); self.assertIsNotNone(b)
                    self.assertEqual(''.join(b.itertext()), ''.join(a.itertext()).replace('612份独立人物档案', str(len(self.actual['people'])) + '份独立人物档案'))
                    self.assertEqual(after, before.replace(b'612&#20221;&#29420;&#31435;', str(len(self.actual['people'])).encode() + b'&#20221;&#29420;&#31435;'))

    def test_live_runtime_and_engines_have_only_reviewed_changes(self):
        for rel in ['site/app.js','site/model.js','site/rosters.js','site/styles.css','site/index.html',
                    'scripts/build.py','scripts/build_rosters.py','scripts/validate_source_contract.py',
                    'scripts/validate_metadata_transition.py','tests/source_date_history.py',
                    'tests/source-date-history.js']:
            expected = (published_root() / rel).read_bytes()
            # Sole approved runtime correction: unknown appointment institutions
            # must not inherit the person's directory affiliation; cache token follows.
            if rel == 'site/app.js':
                before = b"${esc(a.institution||p.department)}"
                after = "${esc(a.institution||'任职机构未核实')}".encode()
                self.assertEqual(expected.count(before), 1)
                expected = expected.replace(before, after)
            # Approved UI-only notice: exact helper and two insertion sites; data unchanged.
            if rel == 'site/app.js':
                before = 'function nationalHome(){'.encode()
                after = 'function dataVersionNotice(){return `<p class="notice data-version">当前已收录 ${(db.people||[]).length} 份人物资料 · ${(db.pending_people||[]).length} 条待完善记录。数据版本：${esc(db.build_id||\'版本未标注\')}。具体来源检索时间见各记录；版本编号不是检索日期或发布时间。</p>`}\nfunction nationalHome(){'.encode()
                self.assertEqual(expected.count(before), 1)
                expected = expected.replace(before, after)
                before = '从学校与院系，走进学者的求学经历与师承关系。</p>'.encode()
                after = '从学校与院系，走进学者的求学经历与师承关系。</p>${dataVersionNotice()}'.encode()
                self.assertEqual(expected.count(before), 1)
                expected = expected.replace(before, after)
                before = '可追溯，比填满每一个空格更重要。</p>'.encode()
                after = '可追溯，比填满每一个空格更重要。</p>${dataVersionNotice()}'.encode()
                self.assertEqual(expected.count(before), 1)
                expected = expected.replace(before, after)
            # Approved representation-only optimization: the actual builder's
            # site atlas writer is compact UTF-8. Every other engine byte stays pinned.
            if rel == 'scripts/build.py':
                before = b"json.dumps(d,ensure_ascii=False,indent=2)+'\\n')"
                after = b"json.dumps(d,ensure_ascii=False,separators=(',',':'),allow_nan=False)+'\\n',encoding='utf-8')"
                self.assertEqual(expected.count(before), 1)
                expected = expected.replace(before, after)
            asset_token_counts = {'site/app.js': 3, 'site/rosters.js': 1, 'site/index.html': 2}
            if rel in asset_token_counts:
                before, after = b'source-dates-20261024', b'data-version-clarity-20261025'
                self.assertEqual(expected.count(before), asset_token_counts[rel])
                expected = expected.replace(before, after)
            actual = (ROOT / rel).read_bytes()
            # Inverse the storage-only delta: every original business byte stays pinned.
            if rel == 'site/app.js':
                added = b"import {loadAtlas} from './atlas-loader.js?v=school-chunks-20261026';\n"
                self.assertEqual(actual.count(added), 1)
                self.assertTrue(actual.startswith(added))
                actual = actual[len(added):]
                after = b"try{db=await loadAtlas();"
                before = "try{const r=await fetch('atlas.json',{cache:'no-store'});if(!r.ok)throw Error('数据文件无法加载');db=await r.json();".encode()
                self.assertEqual(actual.count(after), 1)
                actual = actual.replace(after, before)
            if rel in asset_token_counts:
                self.assertEqual(actual.count(b'school-chunks-20261026'), asset_token_counts[rel])
                actual = actual.replace(b'school-chunks-20261026', b'data-version-clarity-20261025')
            if rel == 'scripts/build.py':
                added = b'from atlas_storage import load_atlas, write_atlas, compute_build_id\n'
                self.assertEqual(actual.count(added), 1)
                actual = actual.replace(added, b'')
                inverse = [
                    (b"d=load_atlas(ROOT/'data/atlas.json')", b"d=json.loads((ROOT/'data/atlas.json').read_text())"),
                    (b"d['build_id']=compute_build_id(ROOT)", b"d['build_id']=hashlib.sha256((ROOT/'data/atlas.json').read_bytes()+(ROOT/'data/rosters.json').read_bytes()+pathlib.Path(__file__).read_bytes()+(ROOT/'scripts/build_rosters.py').read_bytes()+(ROOT/'scripts/deterministic_xlsx.py').read_bytes()).hexdigest()[:16]"),
                    (b"write_atlas(ROOT/'site/atlas.json',d)", b"(ROOT/'site/atlas.json').write_text(json.dumps(d,ensure_ascii=False,separators=(',',':'),allow_nan=False)+'\\n',encoding='utf-8')")]
                for after, before in inverse:
                    self.assertEqual(actual.count(after), 1)
                    actual = actual.replace(after, before)
            self.assertEqual(actual, expected, rel)

    def test_every_affected_workbook_exactly_matches_reviewed_members_and_fields(self):
        from openpyxl import load_workbook
        departments = {p['department_id'] for p in next_mapping()['people'] + seu_mapping()['people'] + scu_mapping()['people'] + uestc_mapping()['people'] + hnu_mapping()['people'] + bnu_mapping()['people'] + tju_mapping()['people'] + nku_mapping()['people'] + dlut_mapping()['people'] + jlu_mapping()['people'] + sdu_mapping()['people'] + csu_mapping()['people'] + hust_mapping()['people'] + whu_mapping()['people'] + xjtu_mapping()['people'] + ruc_mapping()['people'] + ecnu_mapping()['people'] + buaa_mapping()['people'] + tongji_mapping()['people'] + bit_mapping()['people'] + nwpu_mapping()['people'] + scut_mapping()['people']}
        expected_people = self.previous['people'] + self.additions['people']
        for department in departments:
            members = [p for p in expected_people if p['department_id'] == department]
            self.assertEqual(len(members), len({p['person_id'] for p in members}))
            wb = load_workbook(ROOT / 'site/exports' / (department + '.xlsx'), data_only=False)
            try:
                tables = build.department_tables(members, self.actual['updated_at'], self.actual['institutions'])
                self.assertEqual(wb.sheetnames, [name for name, _ in tables])
                for name, rows in tables:
                    expected = [[build.clean(value) if build.clean(value) != '' else None for value in row] for row in rows]
                    actual = [list(row) for row in wb[name].iter_rows(values_only=True)]
                    self.assertEqual(actual, expected, (department, name))
                self.assertEqual([row[0] for row in wb['总览'].iter_rows(min_row=2, values_only=True)], [p['name'] for p in members])
            finally:
                wb.close()

    def test_existing_department_old_rows_survive_member_append(self):
        from openpyxl import load_workbook
        existing = {p['department_id'] for p in self.previous['people']}
        affected = existing & {p['department_id'] for p in uestc_mapping()['people'] + tju_mapping()['people'] + nku_mapping()['people'] + whu_mapping()['people'] + xjtu_mapping()['people'] + ruc_mapping()['people'] + ecnu_mapping()['people'] + buaa_mapping()['people'] + tongji_mapping()['people'] + bit_mapping()['people'] + nwpu_mapping()['people'] + scut_mapping()['people']}
        self.assertEqual(affected, {'uestc-public-administration', 'uestc-applied-math-center', 'tju-math', 'nku-economics', 'whu-math-stat', 'ruc-sociology', 'scut-math'})
        for department in affected:
            old_people = [p for p in self.previous['people'] if p['department_id'] == department]
            for old_person in old_people:
                self.assertEqual(next(p for p in self.actual['people'] if p['person_id'] == old_person['person_id']), old_person)
            wb = load_workbook(ROOT / 'site/exports' / (department + '.xlsx'), data_only=False)
            try:
                for name, rows in build.department_tables(old_people, self.previous['updated_at'], self.previous['institutions']):
                    actual = [list(row) for row in wb[name].iter_rows(min_row=2, values_only=True)]
                    cursor = 0
                    for row in rows[1:]:
                        expected = [build.clean(value) if build.clean(value) != '' else None for value in row]
                        while cursor < len(actual) and actual[cursor] != expected:
                            cursor += 1
                        self.assertLess(cursor, len(actual), (department, name, expected))
                        cursor += 1
            finally:
                wb.close()

    def test_whu_unknown_institution_associations_preserve_full_exported_facts(self):
        from openpyxl import load_workbook
        p = next(p for p in whu_mapping()['people'] if p['person_id'] == 'whu-miaoxiaoping')
        actual = next(x for x in self.actual['people'] if x['person_id'] == p['person_id'])
        self.assertEqual(actual['institutional_associations'], p['institutional_associations'])
        self.assertEqual(len(actual['institutional_associations']), 3)
        self.assertEqual(len(actual['appointments']), 1)
        self.assertEqual(actual['appointments'][0]['title'], '院长')
        wb = load_workbook(ROOT / 'site/exports/whu-public-health.xlsx', data_only=False)
        try:
            rows = [row for row in wb['总览'].iter_rows(values_only=True) if p['person_id'] in row or p['name'] in row]
            self.assertEqual(len(rows), 1)
            text = '\n'.join(str(value) for value in rows[0] if value is not None)
            for association in p['institutional_associations']:
                notes = [note for note in p['unknowns'] if all(str(association[key]) in note for key in ('institution', 'start', 'end'))]
                self.assertEqual(len(notes), 1, association)
                self.assertIn(build.clean(notes[0]), text, association)
        finally:
            wb.close()

    def test_xjtu_source_time_and_role_facts_survive_real_workbook_export(self):
        from openpyxl import load_workbook
        people = {p['person_id']: p for p in xjtu_mapping()['people']}
        for person_id in ('xjtu-yao-xiuying-wave168', 'xjtu-yuyan', 'xjtu-shangyufan-organization'):
            p = people[person_id]
            wb = load_workbook(ROOT / ('site/exports/' + p['department_id'] + '.xlsx'), data_only=False)
            try:
                text = '\n'.join(str(value) for ws in wb for row in ws.iter_rows(values_only=True) for value in row if value is not None)
                self.assertIn(p['name'], text)
                if person_id == 'xjtu-yao-xiuying-wave168':
                    a = p['appointments'][0]
                    self.assertIsNone(a['start'])
                    self.assertEqual(a['university_employment_start'], '2006-07')
                    self.assertIn(build.clean(a['note']), text)
                elif person_id == 'xjtu-yuyan':
                    notes = [n for n in p['unknowns'] if '营养与食品安全系主任' in n]
                    self.assertTrue(notes)
                    for note in notes: self.assertIn(build.clean(note), text)
                else:
                    for year in ('1993', '1996', '2003'):
                        notes = [e['note'] for e in p['education'] if year in e.get('note', '')]
                        self.assertTrue(notes, year)
                        for note in notes: self.assertIn(build.clean(note), text)
            finally:
                wb.close()

    def test_ruc_creative_design_facts_remain_complete_in_actual_overview_export(self):
        from openpyxl import load_workbook
        p = next(p for p in ruc_mapping()['people'] if p['person_id'] == 'ruc-anyaoshun')
        self.assertEqual(len(p['creative_works']), 2)
        self.assertEqual(p['publications'], [])
        wb = load_workbook(ROOT / 'site/exports/ruc-art.xlsx', data_only=False)
        try:
            rows = [row for row in wb['总览'].iter_rows(values_only=True) if p['person_id'] in row or p['name'] in row]
            self.assertEqual(len(rows), 1)
            text = '\n'.join(str(v) for v in rows[0] if v is not None)
            for work in p['creative_works']:
                notes = [note for note in p['unknowns'] if work['title'] in note]
                self.assertEqual(len(notes), 1, work['title'])
                self.assertIn(build.clean(notes[0]), text)
        finally:
            wb.close()

    def test_live_social_workbook_preserves_first_two_source_bound_people(self):
        from openpyxl import load_workbook
        wb = load_workbook(ROOT / 'site/exports/nju-social-sciences.xlsx', data_only=False)
        try:
            text = '\n'.join(str(c.value) for ws in wb for row in ws for c in row if c.value is not None)
            for value in ['施旖旎','杨德睿','2013','2020','未披露']:
                self.assertIn(value, text)
            self.assertNotIn('csu-wangmiao', text)
        finally:
            wb.close()

    def test_buaa_associations_retirement_and_publication_details_remain_visible_in_real_exports(self):
        from openpyxl import load_workbook
        people = {p['person_id']: p for p in buaa_mapping()['people']}
        for person_id in ('buaa-zhou-xuefeng', 'buaa-yuehanwei', 'buaa-jinhao', 'buaa-quanquan'):
            p = people[person_id]
            wb = load_workbook(ROOT / ('site/exports/' + p['department_id'] + '.xlsx'), data_only=False)
            try:
                text = '\n'.join(str(value) for ws in wb for row in ws.iter_rows(values_only=True) for value in row if value is not None)
                self.assertIn(p['name'], text)
                if person_id == 'buaa-zhou-xuefeng':
                    self.assertEqual(len(p['institutional_associations']), 2)
                    self.assertEqual(len(p['appointments']), 2)
                    notes = [n for n in p['unknowns'] if '独立保留不计任职' in n]
                    self.assertEqual(len(notes), 1)
                    for value in ('2004-06', '北京航空航天大学法学院', '1998-07', '2001-04', '山东省高级人民法院'): self.assertIn(value, notes[0])
                    self.assertIn(build.clean(notes[0]), text)
                    for appointment in p['appointments']:
                        self.assertIsNone(appointment['start'])
                        self.assertIn(build.clean(appointment['note']), text)
                elif person_id == 'buaa-yuehanwei':
                    self.assertEqual(p['retirement_year'], 2024)
                    notes = [n for n in p['unknowns'] if '2024' in n]
                    self.assertTrue(notes)
                    for note in notes: self.assertIn(build.clean(note), text)
                    self.assertFalse(p['quick_link_eligible'])
                elif person_id == 'buaa-jinhao':
                    for work in p['publications']: self.assertIn(build.clean(work['note']), text)
                    for value in ('with Chen Xiong', 'with Hewei Shen', 'March 2021', 'January 2020', 'Pages 54-73'): self.assertIn(value, text)
                else:
                    notes = [a['note'] for a in p['appointments'] if '在职信息：在职' in a.get('note', '')]
                    self.assertTrue(notes)
                    for note in notes: self.assertIn(build.clean(note), text)
                    self.assertFalse(p['quick_link_eligible'])
            finally:
                wb.close()

    def test_tongji_composite_degrees_source_dates_and_professional_guidance_survive_real_exports(self):
        from openpyxl import load_workbook
        people = {p['person_id']: p for p in tongji_mapping()['people']}
        for person_id in ('tongji-xielihua', 'tongji-piyong', 'tongji-huominfeng', 'tongji-huding'):
            p = people[person_id]
            wb = load_workbook(ROOT / ('site/exports/' + p['department_id'] + '.xlsx'), data_only=False)
            try:
                text = '\n'.join(str(value) for ws in wb for row in ws.iter_rows(values_only=True) for value in row if value is not None)
                self.assertIn(p['name'], text)
                if person_id == 'tongji-xielihua':
                    self.assertEqual(len(p['education']), 3)
                    self.assertEqual(p['appointments'][0]['start'], '2012-10')
                    self.assertIsNone(p['appointments'][0]['title'])
                    notes = [n for n in p['unknowns'] if '一个教育事件' in n or '不是教授起任日' in n]
                    self.assertEqual(len(notes), 2)
                    for note in notes: self.assertIn(build.clean(note), text)
                elif person_id == 'tongji-piyong':
                    notes = [n for n in p['unknowns'] if '2026-06-01 10:23' in n]
                    self.assertEqual(len(notes), 1)
                    self.assertIn(build.clean(notes[0]), text)
                    for source in p['sources']:
                        if source.get('displayed_date') or source.get('displayed_date_text'):
                            self.assertIsNone(source.get('published_at'))
                elif person_id == 'tongji-huominfeng':
                    self.assertEqual(len(p['mentors']), 1)
                    self.assertEqual(len(p['professional_research_advisors']), 1)
                    relation = p['professional_research_advisors'][0]
                    self.assertFalse(relation['include_in_degree_mentorship_graph'])
                    self.assertIn(build.clean(relation['type_label']), text)
                    self.assertIn(build.clean(relation['note']), text)
                else:
                    notes = [n for n in p['unknowns'] if '1995学士、1999硕士、2013博士获授年' in n]
                    self.assertEqual(len(notes), 1)
                    self.assertIn(build.clean(notes[0]), text)
            finally:
                wb.close()

    def test_bit_source_limits_teaching_and_dates_survive_real_exports(self):
        from openpyxl import load_workbook
        people = {p['person_id']: p for p in bit_mapping()['people']}
        for person_id in ('bit-shifugui', 'bit-shanjiefei', 'bit-sunmiao', 'bit-lumingfei', 'bit-wanglang', 'bit-xiaenjun'):
            p = people[person_id]
            wb = load_workbook(ROOT / ('site/exports/' + p['department_id'] + '.xlsx'), data_only=False)
            try:
                self.assertEqual(len(wb.sheetnames), 5)
                text = '\n'.join(str(value) for ws in wb for row in ws.iter_rows(values_only=True) for value in row if value is not None)
                self.assertIn(p['name'], text)
                for note in p['unknowns']: self.assertIn(build.clean(note), text)
                for work in p['publications']:
                    if work.get('note'): self.assertIn(build.clean(work['note']), text)
                if person_id == 'bit-shifugui':
                    self.assertIsNone(p['publications'][0]['article_number'])
                    self.assertEqual(p['publications'][0]['locator_as_source'], '1190')
                    self.assertIn('1190', text)
                    self.assertIn('退休', text)
                elif person_id == 'bit-shanjiefei':
                    self.assertEqual(p['appointments'][0]['start'], '2013-07')
                    self.assertIsNone(p['appointments'][0]['title'])
                elif person_id == 'bit-sunmiao':
                    self.assertEqual(p['appointments'], [])
                    for value in ('2002', '2019', '设计构成'): self.assertIn(value, text)
                elif person_id == 'bit-lumingfei':
                    self.assertIsNone(p['education'][0]['degree_awarded'])
                    self.assertFalse(p['education'][3]['degree_awarded'])
                    self.assertEqual(len(p['teaching']), 5)
                    for value in ('旁听', '译文', '大学专科'): self.assertIn(value, text)
                elif person_id == 'bit-wanglang':
                    self.assertTrue(all(a['start'] is None for a in p['appointments']))
                    self.assertEqual(p['teaching'][0]['name'], '大学法语')
                    for value in ('2015', '2019', '大学法语'): self.assertIn(value, text)
                else:
                    for value in ('1996-06', '1997-06', '1997-01'): self.assertIn(value, text)
            finally:
                wb.close()

    def test_nwpu_unbound_credentials_dates_guidance_and_creative_limits_survive_real_exports(self):
        from openpyxl import load_workbook
        people = {p['person_id']: p for p in nwpu_mapping()['people']}
        for person_id in ('nwpu-zhaoxueai', 'nwpu-fupengyu', 'nwpu-sunhao-applied-statistics', 'nwpu-sunxiaodi', 'nwpu-zhengye'):
            p = people[person_id]
            wb = load_workbook(ROOT / ('site/exports/' + p['department_id'] + '.xlsx'), data_only=False)
            try:
                self.assertEqual(len(wb.sheetnames), 5)
                text = '\n'.join(str(value) for ws in wb for row in ws.iter_rows(values_only=True) for value in row if value is not None)
                self.assertIn(p['name'], text)
                for note in p['unknowns']: self.assertIn(build.clean(note), text)
                for work in p['publications']:
                    if work.get('note'): self.assertIn(build.clean(work['note']), text)
                if person_id in ('nwpu-zhaoxueai', 'nwpu-fupengyu'):
                    credential = p['credential_labels_as_source']
                    self.assertTrue(credential['degree_awarded'])
                    self.assertIsNone(credential['institution'])
                    self.assertIsNone(credential['date'])
                    self.assertIn('来源报告本人', text)
                if person_id == 'nwpu-zhaoxueai':
                    self.assertEqual(len(p['education']), 3)
                    for value in ('研究生班毕业', '文学学士', '未绑定'): self.assertIn(value, text)
                    self.assertTrue(all(e['degree_awarded'] is None for e in p['education'][:2]))
                elif person_id == 'nwpu-fupengyu':
                    self.assertEqual(len(p['education']), 2)
                    self.assertIsNone(p['education'][0]['degree_awarded'])
                    for value in ('运动人体科学', '北京体育大学', '参编', '全校本科生必修教材'): self.assertIn(value, text)
                elif person_id == 'nwpu-sunhao-applied-statistics':
                    self.assertEqual(p['education'][2]['degree_awarded_date'], '2003-06')
                    self.assertIsNone(p['education'][2]['end'])
                    for value in ('2003-06', '未发生博士后', '赵选民'): self.assertIn(value, text)
                elif person_id == 'nwpu-sunxiaodi':
                    for value in ('终南山', '非论文', '阶段未知', '叶明春'): self.assertIn(value, text)
                    self.assertEqual(p['appointments'], [])
                else:
                    self.assertEqual(p['mentors'][0]['name'], '吴建南')
                    self.assertIn(build.clean(p['mentors'][0]['type_label']), text)
                    self.assertEqual(p['teaching'][0]['role'], '主持课程建设项目')
                    self.assertIn('未读原论文', text)
            finally:
                wb.close()

if __name__ == '__main__': unittest.main()
