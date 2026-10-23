"""Export serialization is stable and preserves worksheet values and links."""
from datetime import datetime
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from build import sheet
from deterministic_xlsx import FIXED_DATETIME, FIXED_ZIP_DATETIME, save_deterministic_xlsx


class DeterministicExportTests(unittest.TestCase):
    def workbook(self, year):
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.created = datetime(year, 2, 3)
        workbook.properties.modified = datetime(year, 3, 4)
        sheet(workbook, '数据', [['名称', '值', '来源'], ['示例', False, 'https://example.test/source'], ['空值', None, '=unsafe()']])
        return workbook

    def test_fixed_metadata_and_repeat_bytes(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory) / name for name in ['first.xlsx', 'second.xlsx']]
            for year, path in zip([2020, 2030], paths):
                save_deterministic_xlsx(self.workbook(year), path)
            self.assertEqual(paths[0].read_bytes(), paths[1].read_bytes())
            with ZipFile(paths[0]) as archive:
                self.assertEqual(archive.namelist(), sorted(archive.namelist()))
                self.assertTrue(all(item.date_time == FIXED_ZIP_DATETIME for item in archive.infolist()))
            workbook = load_workbook(paths[0])
            self.assertEqual(workbook.properties.created, FIXED_DATETIME)
            self.assertEqual(workbook.properties.modified, FIXED_DATETIME)
            self.assertEqual(workbook['数据']['B2'].value, False)
            self.assertIsNone(workbook['数据']['B3'].value)
            self.assertEqual(workbook['数据']['C2'].hyperlink.target, 'https://example.test/source')
            self.assertEqual(workbook['数据']['C3'].value, "'=unsafe()")
            self.assertEqual(workbook['数据'].freeze_panes, 'A2')
            self.assertEqual(workbook['数据'].auto_filter.ref, 'A1:C3')
            workbook.close()


if __name__ == '__main__':
    unittest.main()
