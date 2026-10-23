import unittest,sys,json,copy,pathlib
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]/'scripts'))
from build import validate,clean,ROOT
from openpyxl import load_workbook
class DataTests(unittest.TestCase):
 def setUp(self):self.d=json.loads((ROOT/'data/atlas.json').read_text())
 def test_validate(self):self.assertTrue(validate(self.d))
 def test_evidence_required(self):
  del self.d['people'][0]['education'][0]['source_url']
  with self.assertRaises(AssertionError):validate(self.d)
 def test_duplicate_people(self):
  self.d['people'].append(self.d['people'][0])
  with self.assertRaises(AssertionError):validate(self.d)
 def test_formula_injection(self):self.assertEqual(clean('=HYPERLINK("bad")'),'\'=HYPERLINK("bad")')
 def test_workbooks(self):
  files=list((ROOT/'site/exports').glob('hit*.xlsx'));self.assertEqual(len(files),7)
  for f in files:
   wb=load_workbook(f);self.assertEqual(wb.sheetnames,['总览','经历','导师','论文','来源']);self.assertEqual(wb['总览'].freeze_panes,'A2')
 def test_committee_export_and_conflicted_homepage(self):
  f=next((ROOT/'site/exports').glob('*robot*.xlsx'),None)
  if f is None:
   f=next(f for f in (ROOT/'site/exports').glob('hit*.xlsx') if any(r[0]=='高庆彬' for r in load_workbook(f)['总览'].iter_rows(min_row=2,values_only=True)))
  rows=list(load_workbook(f)['导师'].iter_rows(min_row=2,values_only=True))
  self.assertEqual(len(rows),5)
  self.assertTrue(all(r[7] is None for r in rows))
  self.assertEqual(sum('委员会' in (r[2] or '') for r in rows),4)
 def test_roster_is_separate(self):
  r=json.loads((ROOT/'site/rosters.json').read_text())
  self.assertEqual(len(r['entries']),1381);self.assertEqual(len(self.d['people']),143)
  self.assertTrue(all(e['profile_identity_verified'] is False for e in r['entries'] if e.get('batch_id')=='harbin-002'))
  self.assertEqual(sum(not e['official_profile_url'] for e in r['entries']),163)
 def test_roster_conflicts_preserved(self):
  r=json.loads((ROOT/'site/rosters.json').read_text())
  cs=[c for c in r['conflicts'] if c['type']=='same_url_different_names'];self.assertEqual(len(cs),2)
  for c in cs:self.assertEqual(len([e for e in r['entries'] if e['id'] in c['entry_ids']]),2)
 def test_roster_candidates_not_merged(self):
  r=json.loads((ROOT/'site/rosters.json').read_text());e=next(e for e in r['entries'] if e['name']=='金涛')
  self.assertEqual(e['candidate_person_ids'],['hit-jintao']);self.assertEqual(e['candidate_match_status'],'pending_identity_review')
 def test_roster_exports(self):
  import csv
  w=load_workbook(ROOT/'site/exports/待核对名录.xlsx');self.assertEqual(w.sheetnames,['待核对名录','覆盖说明','分类覆盖','冲突记录','消歧候选','批次来源','排除与待查','原始来源观察','整合依据']);self.assertEqual(w['待核对名录'].max_row,1382)
  with (ROOT/'site/exports/待核对名录.csv').open(encoding='utf-8-sig') as f:self.assertEqual(len(list(csv.reader(f))),1382)
 def test_collaborator_cannot_be_promoted(self):
  p=next(p for p in self.d['people'] if p['name']=='刘冰');p['research_collaborators'][0]['include_in_mentorship_graph']=True
  with self.assertRaises(AssertionError):validate(self.d)
 def test_collaborator_needs_existing_event(self):
  p=next(p for p in self.d['people'] if p['name']=='刘冰');p['research_collaborators'][0]['education_event_id']='missing'
  with self.assertRaises(AssertionError):validate(self.d)
 def test_visiting_collaborator_cannot_bind_degree(self):
  p=next(p for p in self.d['people'] if p['name']=='刘冰');p['research_collaborators'][0]['education_event_id']=p['education'][0]['event_id']
  with self.assertRaises(AssertionError):validate(self.d)
 def test_collaboration_export_is_experience_not_mentor(self):
  from build import department_tables
  people=[p for p in self.d['people'] if p['name'] in ['刘冰','刘大同','张明宇']];tables=dict(department_tables(people,self.d['updated_at']))
  self.assertEqual(len(tables['导师'])-1,sum(len(p['mentors'])+len(p.get('committee_members',[])) for p in people))
  self.assertFalse(any(r[1] in ['Yu Hen Hu','Haitao Liao','张莹莹'] for r in tables['导师'][1:]))
  rows=[r for r in tables['经历'][1:] if r[1] in ['visiting_research_collaborator','research_assistant_collaborator']]
  self.assertEqual(len(rows),3);self.assertTrue(all('非导师认定' in r[14] for r in rows))
 def test_collaboration_export_cannot_join_missing_ids(self):
  from build import department_tables
  p=copy.deepcopy(next(p for p in self.d['people'] if p['name']=='张明宇'));c=p['research_collaborators'][0]
  rows=dict(department_tables([p],self.d['updated_at']))['经历'];row=next(r for r in rows if r[1]=='research_assistant_collaborator')
  self.assertEqual(row[4:6],['2014-03','2014-08']);self.assertIn('清华大学',row[3])
  c.pop('appointment_event_id');rows=dict(department_tables([p],self.d['updated_at']))['经历'];row=next(r for r in rows if r[1]=='research_assistant_collaborator')
  self.assertEqual(row[4:6],[None,None])
if __name__=='__main__':unittest.main()
