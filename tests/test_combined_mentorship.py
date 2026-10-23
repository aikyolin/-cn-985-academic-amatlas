"""Synthetic combined-stage and explicit occupational-event contracts; no public writes."""
import copy
import io
import json
import pathlib
import sys
import unittest
from jsonschema import Draft202012Validator
from openpyxl import Workbook,load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import build

class CombinedMentorshipTests(unittest.TestCase):
 def setUp(self):
  self.d=json.loads((ROOT/'tests/fixtures/multischool.json').read_text());self.p=self.d['people'][1]
  self.p['education']=[dict(event_id='synthetic-stage-'+str(i),degree=degree,event_kind='degree',degree_awarded=None,institution='独立学习机构'+str(i),start='200'+str(i),end='201'+str(i),region='unknown',continent_code='unknown',institution_985='unknown',source_url='https://b.example.test/stage') for i,degree in enumerate(['学士','硕士','博士'])]
  self.m=dict(name='合述导师（虚构）',type='postgraduate_advisor',type_label='硕博合述导师（未拆分阶段）',education_stage='硕博合述，具体分别阶段未拆证',education_event_ids=[e['event_id'] for e in self.p['education'][1:]],stage_scope='combined_masters_doctoral_statement_not_separate_edges',employment_status='unknown',source_url='https://b.example.test/relation',start=None,end=None)
  self.p['mentors']=[self.m];self.schema=Draft202012Validator(json.loads((ROOT/'schema/atlas.schema.json').read_text()))
 def tables(self):return dict(build.department_tables([self.p],self.d['updated_at'],self.d['institutions']))
 def test_combined_validation_and_export_preserve_one_relation_without_changing_evidence(self):
  before=copy.deepcopy(self.d);self.schema.validate(self.d);self.assertTrue(build.validate(self.d));t=self.tables();self.assertEqual(len(t['导师']),2);row=t['导师'][1]
  self.assertEqual(row[2],self.m['type_label']);self.assertEqual(row[3],self.m['education_stage']);self.assertIn(self.m['stage_scope'],row[6]);self.assertIn('合述仅计一条关系，不拆分阶段导师认定',row[6]);self.assertIn('起始时间未披露 至 结束时间未披露',row[6]);self.assertIn(json.dumps(self.m['education_event_ids'],ensure_ascii=False),row[6]);self.assertNotIn('2001',str(row));self.assertNotIn('独立学习机构1',str(row));self.assertEqual(self.d,before)
 def test_real_workbook_roundtrip_retains_both_ids_scope_unknowns_and_one_row(self):
  wb=Workbook();wb.remove(wb.active)
  for name,rows in self.tables().items():build.sheet(wb,name,rows)
  buffer=io.BytesIO();wb.save(buffer);buffer.seek(0);out=load_workbook(buffer);rows=list(out['导师'].iter_rows(values_only=True));self.assertEqual(len(rows),2)
  for value in [self.m['stage_scope'],*self.m['education_event_ids'],'起始时间未披露','结束时间未披露']:self.assertIn(value,rows[1][6])
 def test_bad_plural_types_empty_duplicate_and_nonlocal_ids_reject(self):
  for refs in [None,{},'synthetic-stage-1',[],[''],[' '],[None],['synthetic-stage-1','synthetic-stage-1'],['synthetic-stage-1','missing']]:
   d=copy.deepcopy(self.d);d['people'][1]['mentors'][0]['education_event_ids']=refs
   with self.subTest(refs=refs),self.assertRaises(AssertionError):build.validate(d)
   if refs!=['synthetic-stage-1','missing']:self.assertFalse(self.schema.is_valid(d))
 def test_combined_requires_scope_and_never_accepts_conflicting_single_id(self):
  for key in ['stage_scope','education_stage']:
   d=copy.deepcopy(self.d);del d['people'][1]['mentors'][0][key];self.assertFalse(self.schema.is_valid(d))
   with self.assertRaisesRegex(AssertionError,'explicit stage scope'):build.validate(d)
  self.m['education_event_id']=self.m['education_event_ids'][0]
  with self.assertRaisesRegex(AssertionError,'bindings conflict'):build.validate(self.d)
 def test_single_event_array_or_legacy_singular_remain_valid(self):
  self.m['education_event_ids']=self.m['education_event_ids'][:1];self.m['education_event_id']=self.m['education_event_ids'][0];self.schema.validate(self.d);self.assertTrue(build.validate(self.d));del self.m['education_event_ids'];del self.m['stage_scope'];self.assertTrue(build.validate(self.d));self.schema.validate(self.d)
 def test_explicit_relation_partial_interval_keeps_missing_start(self):
  self.m.update(institution='来源明确合述机构',end='2020-07');row=self.tables()['导师'][1];self.assertIn('来源明确合述机构',row[6]);self.assertIn('起始时间未披露 至 2020-07',row[6]);self.assertNotIn('2001',row[6])

class BoundProfessionalGuidanceTests(unittest.TestCase):
 def setUp(self):
  self.d=json.loads((ROOT/'tests/fixtures/multischool.json').read_text());self.p=self.d['people'][1]
  self.e=dict(event_id='explicit-research',degree=None,event_kind='professional_research',degree_awarded=False,institution='明确科研经历机构',start='2009-08',end='2010-09',region='overseas',continent_code='north_america',institution_985='no',source_url='https://b.example.test/research')
  self.p['education']=[self.e];self.c=dict(name='职业导师（虚构）',type='research_supervisor',relationship_kind='professional_research_advisor',relationship_status='verified_direct_statement',education_event_id=self.e['event_id'],include_in_mentorship_graph=False,include_in_degree_mentorship_graph=False,include_in_verified_mentor_count=False,include_in_verified_professional_guidance_count=True,employment_status='unknown',source_url=self.e['source_url'],start=None,end=None)
  self.p['professional_research_advisors']=[self.c];self.schema=Draft202012Validator(json.loads((ROOT/'schema/atlas.schema.json').read_text()))
 def tables(self):return dict(build.department_tables([self.p],self.d['updated_at'],self.d['institutions']))
 def test_false_award_closed_professional_event_is_valid_and_preserves_own_unknown_dates(self):
  before=copy.deepcopy(self.d);self.schema.validate(self.d);self.assertTrue(build.validate(self.d));row=self.tables()['导师'][1];self.assertEqual(row[2],'职业科研指导（非学位）');self.assertIn('指导起止：起始时间未披露 至 结束时间未披露',row[6]);self.assertIn('对应非学位科研经历：explicit-research',row[6]);self.assertIn('2009-08 至 2010-09（经历区间，不等于指导起止；非博士后认定）',row[6]);self.assertEqual(self.d,before)
 def test_true_unknown_missing_award_and_nonresearch_kinds_reject(self):
  for field,value in [('degree_awarded',True),('degree_awarded',None),('event_kind','degree'),('event_kind','postdoctoral'),('event_kind','visiting_research')]:
   d=copy.deepcopy(self.d);d['people'][1]['education'][0][field]=value
   with self.subTest(field=field,value=value),self.assertRaises(AssertionError):build.validate(d)
   self.assertEqual(build.professional_guidance_event(d['people'][1],d['people'][1]['professional_research_advisors'][0]),{})
  del self.e['degree_awarded']
  with self.assertRaises(AssertionError):build.validate(self.d)
 def test_null_keeps_legacy_unbound_contract_and_never_joins_null_education(self):
  self.c['education_event_id']=None;self.e['event_id']=None;self.schema.validate(self.d);self.assertTrue(build.validate(self.d));row=self.tables()['导师'][1];self.assertNotIn('对应非学位科研经历',row[6]);self.assertNotIn('2009-08',str(row));self.assertNotIn('明确科研经历机构',str(row))
 def test_missing_empty_unknown_foreign_and_ambiguous_event_refs_reject(self):
  for value in ['', ' ', 'missing',17]:
   self.c['education_event_id']=value
   with self.subTest(value=value),self.assertRaises(AssertionError):build.validate(self.d)
  self.c['education_event_id']='explicit-research';self.c['education_event_ids']=['explicit-research']
  with self.assertRaisesRegex(AssertionError,'plural event binding'):build.validate(self.d)
  del self.c['education_event_ids'];self.p['education'].append(copy.deepcopy(self.e))
  with self.assertRaisesRegex(AssertionError,'Duplicate education event'):build.validate(self.d)
 def test_graph_and_count_guards_remain_strict_with_bound_context(self):
  for key,value in [('include_in_mentorship_graph',True),('include_in_degree_mentorship_graph',True),('include_in_verified_mentor_count',True),('include_in_verified_professional_guidance_count',False)]:
   d=copy.deepcopy(self.d);d['people'][1]['professional_research_advisors'][0][key]=value
   with self.subTest(key=key),self.assertRaises(AssertionError):build.validate(d)
   self.assertFalse(self.schema.is_valid(d))
 def test_export_fails_closed_for_invalid_binding_even_without_validation(self):
  for value in [None,'','unknown']:
   self.c['education_event_id']=value;row=self.tables()['导师'][1];self.assertNotIn('2009-08',str(row));self.assertNotIn('明确科研经历机构',str(row))

if __name__=='__main__':unittest.main()
