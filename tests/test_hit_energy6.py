"""Exact HIT six-person additive checks; only the two authorized generated XLSX are read."""
import copy,hashlib,json,os,sys,unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from atlas_storage import load_atlas
import build,build_rosters
from openpyxl import load_workbook
SHA='40bc1424b511fbf76fbddcb664023c1067a3fbd2064a8ae9b6dc70304c41d6f0'
IDS=['hit-wanjie','hit-lanchaohui','hit-keyangguang','hit-zhaowenke','hit-wangwei36-energy','hit-wanglihit-energy']
OLD_IDS=['hit-zhaozhiyuan','hit-yanhan-energy']
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
class HitEnergySixTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.parent=Path(os.environ['HIT_ENERGY6_BASELINE_ROOT']).resolve()
  cls.receipt_path=Path(os.environ['HIT_ENERGY6_BUILD_RECEIPT']).resolve()
  if cls.parent==ROOT or cls.receipt_path.is_relative_to(ROOT):raise ValueError('separate immutable baseline/external receipt required')
  fp=ROOT/'tests/fixtures/hit-energy6-source-mapping.json'
  if sha(fp)!=SHA:raise ValueError('fixed mapping SHA mismatch')
  cls.fixed=json.loads(fp.read_text());cls.base=load_atlas(cls.parent/'data/atlas.json');cls.actual=load_atlas(ROOT/'data/atlas.json');cls.site=load_atlas(ROOT/'site/atlas.json');cls.receipt=json.loads(cls.receipt_path.read_text())
 def test_01_old_1151_whole_objects_and_all_other_root_values_are_preserved(self):
  self.assertEqual(hashlib.sha256(json.dumps(self.base,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest(),'b6a77b6aa4d88a89a6c4558587477058e9f625c900d44fc134ad06b0337cac13');self.assertEqual(len(self.base['people']),1151);self.assertEqual(len(self.actual['people']),1157)
  self.assertEqual(self.actual['people'][:1151],self.base['people']);self.assertEqual(self.actual['people'][1151:],self.fixed['people'])
  self.assertEqual([p['person_id'] for p in self.fixed['people']],IDS)
  self.assertEqual(self.actual['profile_batches'],self.base['profile_batches']+self.fixed['profile_batches'])
  restored=copy.deepcopy(self.actual);restored['people']=restored['people'][:1151];restored['profile_batches']=restored['profile_batches'][:-1]
  self.assertEqual(restored,self.base);self.assertEqual(len(self.actual['pending_people']),84)
  self.assertEqual((ROOT/'data/rosters.json').read_bytes(),(self.parent/'data/rosters.json').read_bytes())
  for k,v in self.actual.items():self.assertEqual(self.site[k],v,k)
  others=[p for p in (ROOT/'site/people').glob('*.json') if p.name!='hit.json'];self.assertEqual(len(others),38)
  for p in others:self.assertEqual(p.read_bytes(),(self.parent/'site/people'/p.name).read_bytes(),p.name)
 def test_02_source_semantics_and_negative_projection_changes(self):
  people=self.actual['people'][1151:];events=[e for p in people for e in p['education']]
  self.assertEqual(len(events),14);self.assertEqual([sum(e['degree_awarded'] is v for e in events)for v in (True,None,False)],[3,9,2])
  self.assertEqual(sum(len(p['appointments'])for p in people),20);self.assertEqual(sum(len(p['mentors'])for p in people),2);self.assertEqual(sum(len(p['publications'])for p in people),12)
  for p in people:
   self.assertEqual(p['school_id'],'hit');self.assertEqual(p['department_id'],'hit-energy');self.assertEqual(p['campus_ids'],['harbin'])
   for e in p['education']:self.assertIn('degree_awarded_date',e);self.assertIsNone(e['degree_awarded_date'])
   for m in p['mentors']:self.assertEqual(m['employment_status'],'unknown');self.assertIsNone(m['person_id']);self.assertIsNone(m['formal_dissertation_advisor_verified'])
  self.assertEqual(people[5]['education'],[])
  changed=copy.deepcopy(people);changed[0]['education'][0]['degree_awarded']=True
  self.assertNotEqual(changed,self.fixed['people'])
  changed=copy.deepcopy(people);changed[2]['mentors'][0]['education_event_id']='hit-keyangguang-education-1'
  self.assertNotEqual(changed,self.fixed['people'])
  changed=copy.deepcopy(people);changed[4]['publications'][1]['year']=2026
  self.assertNotEqual(changed,self.fixed['people'])
 def assert_tables(self,path,tables):
  book=load_workbook(path,read_only=True,data_only=False)
  try:
   self.assertEqual(book.sheetnames,list(tables))
   for name,rows in tables.items():
    expected=[[build.clean(v) if build.clean(v)!='' else None for v in row]for row in rows]
    width=max(map(len,expected),default=0)
    expected=[row+[None]*(width-len(row)) for row in expected]
    actual=[list(row)for row in book[name].iter_rows(values_only=True)]
    self.assertEqual(actual,expected,name)
  finally:book.close()
 def test_03_only_hit_energy_five_tables_and_global_roster_nine_tables(self):
  members=[p for p in self.actual['people']if p['department_id']=='hit-energy'];self.assertEqual([p['person_id']for p in members],OLD_IDS+IDS)
  old_members=[p for p in self.base['people']if p['department_id']=='hit-energy'];self.assertEqual([p['person_id']for p in old_members],OLD_IDS);self.assertEqual(members[:2],old_members)
  tables=dict(build.department_tables(members,self.actual['updated_at'],self.actual['institutions']));self.assertEqual(list(tables),['总览','经历','导师','论文','来源'])
  self.assert_tables(ROOT/'site/exports/hit-energy.xlsx',tables)
  rosters=json.loads((ROOT/'data/rosters.json').read_text());rt=build_rosters.roster_tables(build_rosters.prepare_rosters(rosters,self.actual),self.actual);self.assertEqual(len(rt),9)
  self.assert_tables(ROOT/'site/exports/待核对名录.xlsx',rt)
  # These exact source-derived checks do not rely solely on shared table functions.
  self.assertEqual([r[0]for r in tables['总览'][1:]],[p['name']for p in members]);self.assertEqual(len(tables['导师'])-1,5);self.assertEqual(len(tables['论文'])-1,14)
  headers=tables['经历'][0];self.assertEqual(headers[11],'来源记载学位');self.assertEqual(headers[14],'备注及冲突')
  self.assertEqual(len(tables['经历'])-1,46)
  for p in members:
   rows=[r for r in tables['经历'][1:]if r[0]==p['name']];self.assertEqual(len(rows),len(p['education'])+len(p['appointments'])+len(p.get('research_collaborators',[])))
   for row,e in zip(rows[:len(p['education'])],p['education'],strict=True):
    self.assertEqual(row[headers.index('事件类型')],e['event_kind']);self.assertEqual(row[headers.index('阶段/职位')],e['degree']);self.assertIs(row[headers.index('来源记载学位')],e['degree_awarded']);self.assertEqual(row[headers.index('开始')],e['start']);self.assertEqual(row[headers.index('结束')],e['end']);self.assertEqual(row[headers.index('备注及冲突')],e['note'])
   for row,a in zip(rows[len(p['education']):len(p['education'])+len(p['appointments'])],p['appointments'],strict=True):
    self.assertEqual(row[headers.index('事件类型')],'employment');self.assertEqual(row[headers.index('阶段/职位')],a['title']);self.assertEqual(row[headers.index('开始')],a['start']);self.assertEqual(row[headers.index('结束')],a['end']);self.assertEqual(row[headers.index('备注及冲突')],a['note']);self.assertIsNone(row[headers.index('来源记载学位')])
  lan=next(p for p in members if p['person_id']=='hit-lanchaohui');c=lan['research_collaborators'][0]
  self.assertEqual(c['name'],'Igor Kaganovich');self.assertEqual(c['appointment_event_id'],'hit-lanchaohui-career-4');self.assertFalse(c['include_in_mentor_count'])
  row=next(r for r in tables['经历'][1:]if r[0]==lan['name'] and r[1]=='source_reported_collaborator')
  self.assertIn(c['name'],row[14]);self.assertIn('非导师认定',row[14]);self.assertEqual(row[4:6],['2017-03','2019-03']);self.assertIsNone(row[11])
 def test_04_receipt_exact_two_workbooks_and_complete_current_output_hashes(self):
  r=self.receipt;self.assertEqual(r['status'],'HIT_ENERGY6_RESTRICTED_BUILD_PASS');self.assertIsNone(r['base_commit']);self.assertEqual(r['build_id'],self.site['build_id'])
  workbooks={'site/exports/hit-energy.xlsx','site/exports/待核对名录.xlsx'}
  outputs=workbooks|{'site/people/'+p.name for p in (ROOT/'data/people').glob('*.json')}|{'site/atlas.json','site/rosters.json','site/exports/待核对名录.csv'}
  self.assertEqual(len(outputs),44);self.assertEqual(set(r['generated_workbooks']),workbooks);self.assertEqual(set(r['generated_output_inventory']),outputs);self.assertEqual(set(r['generated_sha256']),outputs)
  for rel in outputs:self.assertEqual(sha(ROOT/rel),r['generated_sha256'][rel],rel)
  self.assertEqual(r['old_full_suite'],'NOT_RUN');self.assertFalse(any('xmu-'in p for p in workbooks))
if __name__=='__main__':unittest.main()
